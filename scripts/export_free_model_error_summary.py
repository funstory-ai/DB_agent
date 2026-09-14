"""Export free-model track_translate_page_data failure summary to Excel."""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parent))
from db import query

OUT = Path(__file__).resolve().parents[1] / "output"
OUT.mkdir(exist_ok=True)
OUT_FILE = OUT / "free_model_translate_error_summary_20260827.xlsx"

TIME_START = "2026-08-20"
TIME_END = "2026-08-27"

OVERVIEW_SQL = f"""
SELECT metric, value FROM (
  SELECT '事件名' AS metric, 'track_translate_page_data' AS value, 1 AS ord
  UNION ALL SELECT '翻译服务(ex_char_arg1)', 'free-model', 2
  UNION ALL SELECT '结果筛选(ex_char_arg2)', 'failure', 3
  UNION ALL SELECT '时间窗(北京日)', '{TIME_START} ~ 2026-08-26', 4
  UNION ALL SELECT '查数日期', '2026-08-27', 5
  UNION ALL SELECT 'failure 总行数', cast(count(*) AS varchar), 6
  FROM immdata.user_behaviour
  WHERE created_ts >= '{TIME_START}' AND created_ts < '{TIME_END}'
    AND event_name = 'track_translate_page_data'
    AND ex_char_arg1 = 'free-model' AND ex_char_arg2 = 'failure'
  UNION ALL SELECT 'free-model 全结果行数', cast(count(*) AS varchar), 7
  FROM immdata.user_behaviour
  WHERE created_ts >= '{TIME_START}' AND created_ts < '{TIME_END}'
    AND event_name = 'track_translate_page_data'
    AND ex_char_arg1 = 'free-model'
) t ORDER BY ord
"""

BY_RESULT_SQL = f"""
SELECT ex_char_arg2 AS first_visible_result, count(*) AS cnt,
  round(count(*) * 100.0 / sum(count(*)) OVER (), 4) AS pct
FROM immdata.user_behaviour
WHERE created_ts >= '{TIME_START}' AND created_ts < '{TIME_END}'
  AND event_name = 'track_translate_page_data' AND ex_char_arg1 = 'free-model'
GROUP BY 1 ORDER BY cnt DESC
"""

BUCKET_SQL = f"""
WITH base AS (
  SELECT ex_char_arg4 AS failure_code, ex_char_arg5 AS failure_reason
  FROM immdata.user_behaviour
  WHERE created_ts >= '{TIME_START}' AND created_ts < '{TIME_END}'
    AND event_name = 'track_translate_page_data'
    AND ex_char_arg1 = 'free-model' AND ex_char_arg2 = 'failure'
), tot AS (SELECT count(*) AS n FROM base)
SELECT
  CASE
    WHEN failure_code = 'network' OR failure_reason = 'network_error' THEN '网络(network_error)'
    WHEN failure_reason = 'render_error' THEN '客户端渲染(render_error)'
    WHEN failure_reason = 'provider_error' OR failure_code LIKE 'http_5%' THEN '服务商5xx(provider_error)'
    WHEN failure_reason = 'quota_or_rate_limit' OR failure_code = 'http_429' THEN '限流429'
    WHEN failure_reason = 'auth_error' OR failure_code IN ('http_401','http_403') THEN '鉴权401/403'
    WHEN failure_reason = 'invalid_response' AND failure_code = 'translation' THEN '无HTTP状态/响应异常(invalid_response)'
    WHEN failure_reason = 'invalid_response' AND failure_code LIKE 'http_4%' THEN 'HTTP4xx(非429/401/403)'
    ELSE '其它'
  END AS bucket,
  CASE
    WHEN failure_code = 'network' OR failure_reason = 'network_error' THEN '国内端点兜底'
    WHEN failure_reason = 'render_error' THEN '客户端问题(非API)'
    WHEN failure_reason = 'provider_error' OR failure_code LIKE 'http_5%' THEN '换provider/谷歌兜底'
    WHEN failure_reason = 'quota_or_rate_limit' OR failure_code = 'http_429' THEN '扩容/降频'
    WHEN failure_reason = 'auth_error' OR failure_code IN ('http_401','http_403') THEN '鉴权配置'
    WHEN failure_reason = 'invalid_response' AND failure_code = 'translation' THEN '查服务返回/模型/provider'
    WHEN failure_reason = 'invalid_response' AND failure_code LIKE 'http_4%' THEN '查请求参数/路由/术语传参'
    ELSE '个案排查'
  END AS action_hint,
  count(*) AS cnt,
  round(count(*) * 100.0 / (SELECT n FROM tot), 4) AS pct_of_failures
FROM base GROUP BY 1, 2 ORDER BY cnt DESC
"""

DETAIL_SQL = f"""
WITH base AS (
  SELECT ex_char_arg4 AS failure_code, ex_char_arg5 AS failure_reason, ex_char_arg6 AS failure_stage
  FROM immdata.user_behaviour
  WHERE created_ts >= '{TIME_START}' AND created_ts < '{TIME_END}'
    AND event_name = 'track_translate_page_data'
    AND ex_char_arg1 = 'free-model' AND ex_char_arg2 = 'failure'
), tot AS (SELECT count(*) AS n FROM base)
SELECT failure_code, failure_reason, failure_stage, count(*) AS cnt,
  round(count(*) * 100.0 / (SELECT n FROM tot), 4) AS pct_of_failures
FROM base GROUP BY 1, 2, 3 ORDER BY cnt DESC
"""

NOTES = [
    ("说明", "内容"),
    ("事件语义", "网页翻译会话「首个可见译文」结果；非错误弹窗专用事件"),
    ("failure_code", "ex_char_arg4；如 network / translation / http_502"),
    ("failure_reason", "ex_char_arg5；如 network_error / invalid_response / provider_error"),
    ("failure_stage", "ex_char_arg6；request / response / render"),
    ("与弹窗关系", "首屏失败即上报，不等同 error_modal_show；段落错误需累计才弹窗"),
    ("插件依据", "F:\\extension src/analytics/first_visible_translation.ts"),
    ("SQL", "sql/reports/query_free_translate_error_tracking_20260827.sql"),
]


def write_sheet(wb: Workbook, title: str, cols: list[str], rows: list[tuple], header: bool = True):
    ws = wb.create_sheet(title=title)
    if header:
        ws.append(cols)
        for cell in ws[1]:
            cell.font = Font(bold=True)
    for row in rows:
        ws.append(list(row))
    for col in ws.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, 10), 48)


def main():
    wb = Workbook()
    wb.remove(wb.active)

    _, overview = query(OVERVIEW_SQL)
    write_sheet(wb, "overview", ["metric", "value"], overview)

    _, by_result = query(BY_RESULT_SQL)
    write_sheet(
        wb,
        "by_result",
        ["first_visible_result", "cnt", "pct"],
        by_result,
    )

    _, bucket = query(BUCKET_SQL)
    write_sheet(
        wb,
        "bucket_summary",
        ["bucket", "action_hint", "cnt", "pct_of_failures"],
        bucket,
    )

    _, detail = query(DETAIL_SQL)
    write_sheet(
        wb,
        "detail_code_reason",
        ["failure_code", "failure_reason", "failure_stage", "cnt", "pct_of_failures"],
        detail,
    )

    write_sheet(wb, "notes", NOTES[0], NOTES[1:])

    wb.save(OUT_FILE)
    print(f"Wrote {OUT_FILE}")


if __name__ == "__main__":
    main()
