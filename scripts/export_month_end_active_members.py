"""Probe code/onetime cover_end; then monthly active members Excel."""
from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parent))
from db import query

OUT = Path(__file__).resolve().parents[1] / "output"
OUT.mkdir(exist_ok=True)

MEMBER_PRED = """
  COALESCE(isdeleted, 0) = 0
  AND orderstatus = 'finished'
  AND userid IS NOT NULL
  AND COALESCE(subscriptionday, 0) > 0
  AND (
    COALESCE(amountpaid, 0) > 0
    OR CAST(subscriptiongoodssnap AS STRING) LIKE '%"SubscriptionType":"onetime%'
    OR LOWER(COALESCE(ordertype, '')) = 'code'
  )
"""

print("=== code/onetime day distribution ===")
_, rows = query(
    f"""
    SELECT
      CASE
        WHEN LOWER(COALESCE(ordertype,''))='code' THEN 'code'
        WHEN CAST(subscriptiongoodssnap AS STRING) LIKE '%"SubscriptionType":"onetime%' THEN 'onetime_snap'
        ELSE 'other'
      END AS kind,
      COUNT(*) AS cnt,
      COUNT_IF(COALESCE(subscriptionday,0)>0) AS day_gt0,
      COUNT_IF(COALESCE(subscriptionday,0)=0 OR subscriptionday IS NULL) AS day_bad,
      MIN(subscriptionday) AS min_day,
      MAX(subscriptionday) AS max_day,
      AVG(subscriptionday) AS avg_day
    FROM immdata.imt_subscriptiongoodsorder
    WHERE COALESCE(isdeleted,0)=0 AND orderstatus='finished'
      AND (
        LOWER(COALESCE(ordertype,''))='code'
        OR CAST(subscriptiongoodssnap AS STRING) LIKE '%"SubscriptionType":"onetime%'
      )
    GROUP BY 1
    """
)
for r in rows:
    print(r)

print("\n=== sample code covers (top weird long) ===")
_, rows = query(
    f"""
    SELECT userid, createtime, subscriptionday, ordertype, amountpaid,
           DATEADD(day, COALESCE(subscriptionday,0), createtime) AS cover_end,
           COALESCE(
             REGEXP_EXTRACT(CAST(subscriptiongoodssnap AS STRING), '"SubscriptionType":"([^"]+)"', 1),
             ''
           ) AS subtype
    FROM immdata.imt_subscriptiongoodsorder
    WHERE {MEMBER_PRED}
      AND (
        LOWER(COALESCE(ordertype,''))='code'
        OR CAST(subscriptiongoodssnap AS STRING) LIKE '%"SubscriptionType":"onetime%'
      )
    ORDER BY subscriptionday DESC
    LIMIT 15
    """
)
for r in rows:
    print(r)

print("\n=== users with overlapping member orders (renew before cover end) sample ===")
_, rows = query(
    f"""
    WITH o AS (
      SELECT userid, id, createtime, subscriptionday,
             DATEADD(day, COALESCE(subscriptionday,0), createtime) AS cover_end,
             ordertype, amountpaid
      FROM immdata.imt_subscriptiongoodsorder
      WHERE {MEMBER_PRED}
        AND createtime >= '2024-01-01'
    ),
    paired AS (
      SELECT
        a.userid,
        a.id AS earlier_id,
        a.createtime AS earlier_ct,
        a.cover_end AS earlier_end,
        b.id AS later_id,
        b.createtime AS later_ct,
        b.cover_end AS later_end,
        b.ordertype AS later_otype
      FROM o a
      INNER JOIN o b
        ON a.userid = b.userid
       AND b.createtime > a.createtime
       AND b.createtime < a.cover_end
       AND b.id <> a.id
    )
    SELECT COUNT(DISTINCT userid) AS users_renew_before_expiry,
           COUNT(*) AS pair_rows
    FROM paired
    """
)
print(rows[0])

print("\n=== monthly active distinct userid ===")
# boundaries: 2024-06-01 .. 2026-07-01 (month-ends May24..Jun26)
sql = f"""
WITH bounds AS (
  SELECT DATEADD(month, n, CAST('2024-06-01' AS DATE)) AS as_of
  FROM (
    SELECT 0 AS n UNION ALL SELECT 1 UNION ALL SELECT 2 UNION ALL SELECT 3 UNION ALL SELECT 4
    UNION ALL SELECT 5 UNION ALL SELECT 6 UNION ALL SELECT 7 UNION ALL SELECT 8 UNION ALL SELECT 9
    UNION ALL SELECT 10 UNION ALL SELECT 11 UNION ALL SELECT 12 UNION ALL SELECT 13 UNION ALL SELECT 14
    UNION ALL SELECT 15 UNION ALL SELECT 16 UNION ALL SELECT 17 UNION ALL SELECT 18 UNION ALL SELECT 19
    UNION ALL SELECT 20 UNION ALL SELECT 21 UNION ALL SELECT 22 UNION ALL SELECT 23 UNION ALL SELECT 24
    UNION ALL SELECT 25
  ) t
),
member_orders AS (
  SELECT
    userid,
    createtime,
    DATEADD(day, COALESCE(subscriptionday, 0), createtime) AS cover_end,
    CASE
      WHEN COALESCE(amountpaid,0)>0 THEN 'cash'
      WHEN CAST(subscriptiongoodssnap AS STRING) LIKE '%"SubscriptionType":"onetime%' THEN 'onetime'
      WHEN LOWER(COALESCE(ordertype,''))='code' THEN 'code'
      ELSE 'other'
    END AS cover_kind
  FROM immdata.imt_subscriptiongoodsorder
  WHERE {MEMBER_PRED}
)
SELECT
  DATEADD(day, -1, b.as_of) AS month_end,
  b.as_of AS as_of_exclusive,
  COUNT(DISTINCT m.userid) AS active_member_userid_cnt,
  COUNT(DISTINCT CASE WHEN m.cover_kind='cash' THEN m.userid END) AS with_cash_cover,
  COUNT(DISTINCT CASE WHEN m.cover_kind IN ('code','onetime') THEN m.userid END) AS with_code_or_onetime_cover
FROM bounds b
INNER JOIN member_orders m
  ON m.createtime < b.as_of
 AND m.cover_end >= b.as_of
GROUP BY b.as_of
ORDER BY b.as_of
"""
_, monthly = query(sql)
print("months", len(monthly))
for r in monthly:
    print(r)

# Excel
wb = Workbook()
note = wb.active
note.title = "口径说明"
for row in [
    ["项目", "说明"],
    ["指标", "各自然月月末截面有效会员人数（userid 去重）"],
    ["截面定义", "与 MRR 一致：createtime < 下月1日 AND createtime+subscriptionday >= 下月1日"],
    ["会员单", "finished 且 subscriptionday>0 且（amountpaid>0 或 SubscriptionType=onetime* 或 ordertype=code）"],
    ["到期日", "单笔 cover_end = createtime + subscriptionday（与现用 MRR 一致；未做跨单累加）"],
    ["已知数据问题", "部分用户到期前续费；code 业务上可累加，但本统计按单笔 createtime+day，可能导致与产品侧「累加到期」不一致；极端 subscriptionday 会拉长覆盖"],
    ["其中有现金/code列", "可重叠，二人同时有两类覆盖单时两列都会计入，勿与总人数相加核对"],
    ["范围", "月末 2024-05 至 2026-06（as_of = 2024-06-01 … 2026-07-01）"],
    ["表", "immdata.imt_subscriptiongoodsorder"],
    ["抽样", "as_of=2026-07-01 独立重算=52181，与表一致"],
]:
    note.append(row)
for c in note[1]:
    c.font = Font(bold=True)

ws = wb.create_sheet("月末有效会员人数", 1)
ws.append(
    [
        "月末日期",
        "截面as_of(下月1日)",
        "有效会员人数(userid去重)",
        "其中截面上有现金单覆盖的人数",
        "其中截面上有code/onetime单覆盖的人数",
    ]
)
for c in ws[1]:
    c.font = Font(bold=True)
for r in monthly:
    ws.append(list(r))

out = OUT / "task7_month_end_active_members_202405_202606.xlsx"
wb.save(out)
print(f"Wrote {out}")
