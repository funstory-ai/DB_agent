# -*- coding: utf-8 -*-
"""T018 asof=2026-09-28 · 剔老号（1天缓冲）全量周表 → Excel（不改正式 asof 归档）。"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parent))
import export_t018_asof_20260928 as base

ROOT = Path(__file__).resolve().parents[1]
base.SQL_PATH = (
    ROOT / "sql" / "reports" / "query_paywall_ab_longterm_excl_old_acct_20260928.sql"
)
base.OUT_XLSX = (
    ROOT
    / "output"
    / "paywall_ab"
    / "paywall_ab_longterm_excl_old_acct_20260928.xlsx"
)


def retitle(path: Path) -> None:
    wb = load_workbook(path)
    bold = Font(bold=True)
    ws = wb["转化总览"]
    ws["A1"] = "付费墙 A/B/C/D · 剔老号版（1 天缓冲）"
    ws["A1"].font = bold
    ws["A2"] = (
        f"asof={base.ASOF}；分母=P034 激活且剔除老号："
        "userid.createtime < 激活 first_ts − 1 天"
        "（userid=deviceinfo ∪ 激活行 ex_int_arg1）。"
        "非正式主口径；全量对照见 paywall_ab_longterm_activation_20260928.xlsx。"
    )
    if "口径与补算" in wb.sheetnames:
        ws2 = wb["口径与补算"]
        ws2["A1"] = "口径（剔老号版）"
        ws2["A1"].font = bold
        # 在说明区追加几行
        ws2["A20"] = "老号剔除"
        ws2["B20"] = (
            "imt_userinfo.createtime < first_ts - 1 天；"
            "缓冲避免同会话/入表延迟误伤；不用 deviceinfo 时间判老号"
        )
        ws2["A21"] = "正式口径"
        ws2["B21"] = "仍以 asof_20260928 全量激活版为准；本文件为影响评估交付"
    wb.save(path)
    print(f"retitled {path}", flush=True)


def main() -> None:
    rows = base.run()
    out = base.write_xlsx(rows)
    retitle(out)
    # 人数应小于全量：与影响评估对照
    agg = {r["exp_group"]: r for r in rows if r["slice_type"] == "aggregate"}
    print("\n汇总（剔老号）", flush=True)
    for g in "ABCD":
        a = agg[g]
        print(
            g,
            "n=", a["new_devices"],
            "reg=", a["register_rate_pct"],
            "paid=", a["paid_to_asof_pct"],
            "w1=", a["ret_w1_pct"],
            flush=True,
        )
    base.sample_validate(rows)


if __name__ == "__main__":
    main()
