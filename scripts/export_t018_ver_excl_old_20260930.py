# -*- coding: utf-8 -*-
"""T048: T018 asof=09-28 · 语义化版本≥1.32.6 + 剔老号 → Excel。"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parent))
import export_t018_asof_20260928 as base

ROOT = Path(__file__).resolve().parents[1]
base.SQL_PATH = (
    ROOT / "sql" / "reports" / "query_paywall_ab_longterm_ver_excl_old_20260930.sql"
)
base.OUT_XLSX = (
    ROOT
    / "output"
    / "paywall_ab"
    / "paywall_ab_longterm_ver1326_excl_old_20260930.xlsx"
)


def retitle(path: Path) -> None:
    wb = load_workbook(path)
    bold = Font(bold=True)
    ws = wb["转化总览"]
    ws["A1"] = "付费墙 A/B/C/D · 版本≥1.32.6（语义化）+ 剔老号"
    ws["A1"].font = bold
    ws["A2"] = (
        f"asof={base.ASOF}；P040 版本=实验起始 1.32.6（major/minor/patch，"
        "含 1.32.10+；禁止字符串 >=）；P039 剔老号 createtime < first_ts−1天。"
        "对照全量字符串版：paywall_ab_longterm_activation_20260928.xlsx。"
    )
    if "口径与补算" in wb.sheetnames:
        ws2 = wb["口径与补算"]
        ws2["A1"] = "口径（P040 版本 + P039 剔老号）"
        ws2["A1"].font = bold
        ws2["A20"] = "版本（P040）"
        ws2["B20"] = (
            "语义化 ≥1.32.6；修复旧滤掉 1.32.10；"
            "老客户端带 ab_tag 但不在实验 UX 内须踢出"
        )
        ws2["A21"] = "老号（P039）"
        ws2["B21"] = "imt_userinfo.createtime < first_ts - 1 天"
        ws2["A22"] = "需求"
        ws2["B22"] = "T048；关联 T018"
    wb.save(path)
    print(f"retitled {path}", flush=True)


def main() -> None:
    rows = base.run()
    out = base.write_xlsx(rows)
    retitle(out)
    print("\n=== 累计付费% by week (ver+excl_old) ===", flush=True)
    weeks = [
        "natural_week_0824_0830",
        "natural_week_0831_0906",
        "natural_week_0907_0913",
        "natural_week_0914_0920",
        "natural_week_0921_0927",
        "aggregate_0821_0928",
    ]
    print(f"{'slice':28} {'A':>7} {'B':>7} {'C':>7} {'D':>7}", flush=True)
    for sl in weeks:
        vals = {}
        ns = {}
        for g in "ABCD":
            r = next(x for x in rows if x["slice"] == sl and x["exp_group"] == g)
            vals[g] = float(r["paid_to_asof_pct"])
            ns[g] = int(r["new_devices"])
        print(
            f"{sl:28} {vals['A']:6.3f}% {vals['B']:6.3f}% "
            f"{vals['C']:6.3f}% {vals['D']:6.3f}%",
            flush=True,
        )
        if sl == "aggregate_0821_0928":
            print(
                f"{'n_agg':28} {ns['A']:7d} {ns['B']:7d} {ns['C']:7d} {ns['D']:7d}",
                flush=True,
            )
    base.sample_validate(rows)


if __name__ == "__main__":
    main()
