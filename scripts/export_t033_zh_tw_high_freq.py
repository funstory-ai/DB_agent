# -*- coding: utf-8 -*-
"""T033: 繁中高频（免费/退订/试用未付费，HF=近90天>300）+ 即将到期月费 → 四 sheet Excel。"""
from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from db import query
from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
SQL_PATH = ROOT / "sql" / "reports" / "query_zh_hant_high_freq_user_20260910.sql"
OUT_DIR = ROOT / "output"
OUT_DIR.mkdir(exist_ok=True)

SHEETS = {
    "退订": [
        ("userid", "userid"),
        ("email", "邮箱"),
        ("translate_page_90d", "近90天translate_page次数"),
        ("last_cover_end", "会员覆盖到期时间"),
        ("days_since_expiry", "已过期天数"),
    ],
    "试用未付费": [
        ("userid", "userid"),
        ("email", "邮箱"),
        ("translate_page_90d", "近90天translate_page次数"),
        ("first_trial_ts_90d", "近90天首次试用时间"),
    ],
    "免费": [
        ("userid", "userid"),
        ("email", "邮箱"),
        ("translate_page_90d", "近90天translate_page次数"),
    ],
    "即将到期月费": [
        ("userid", "userid"),
        ("email", "邮箱"),
        ("last_cover_end", "预计到期时间"),
        ("days_to_expiry", "距到期天数"),
    ],
}


def main() -> None:
    sql = SQL_PATH.read_text(encoding="utf-8")
    print("running T033 SQL (HF=90d)...", flush=True)
    cols, rows = query(sql)
    print("cols", cols, "nrows", len(rows), flush=True)
    by = {k: [] for k in SHEETS}
    col_idx = {c: i for i, c in enumerate(cols)}
    for r in rows:
        label = r[col_idx["cohort_label"]]
        if label in by:
            by[label].append(r)

    stamp = datetime.now().strftime("%Y%m%d")
    out = OUT_DIR / f"zh_tw_high_freq_user_hf90d_{stamp}.xlsx"
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    summary = wb.create_sheet("汇总", 0)
    summary.append(["人群", "人数", "备注"])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    summary.append(["口径", "", "高频=近90天translate_page>300；付费窗不变"])

    for name, fields in SHEETS.items():
        ws = wb.create_sheet(name)
        headers = [h for _, h in fields]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        data = by.get(name, [])
        note = "不筛高频" if name == "即将到期月费" else "HF90d>300"
        summary.append([name, len(data), note])
        for r in data:
            ws.append([r[col_idx[k]] for k, _ in fields])
        print(f"sheet {name}: {len(data)}", flush=True)

    wb.save(out)
    print("wrote", out, flush=True)


if __name__ == "__main__":
    main()
