# -*- coding: utf-8 -*-
"""T043: 近1年曾 MAX 用户 + 母语 + 仍在MAX/已退订 → Excel。"""
from __future__ import annotations

import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from db import query
from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[1]
SQL_PATH = ROOT / "sql" / "reports" / "query_max_users_lang_1y_20260924.sql"
OUT_DIR = ROOT / "output"
OUT_DIR.mkdir(exist_ok=True)

HEADERS = [
    ("userid", "userid"),
    ("email", "邮箱"),
    ("native_language", "账号母语"),
    ("max_status", "MAX状态"),
    ("max_cover_end", "MAX覆盖到期时间"),
    ("first_max_ts", "首次MAX订单时间"),
    ("last_max_ts", "最近MAX订单时间"),
    ("max_order_cnt", "MAX订单数"),
    ("has_paid_max", "是否有现金MAX"),
    ("has_code_max", "是否有兑换码MAX"),
    ("last_max_subscriptiontype", "最近MAX订阅类型"),
    ("last_max_ordertype", "最近MAX订单渠道"),
]


def main() -> None:
    sql = SQL_PATH.read_text(encoding="utf-8")
    print("running T043 SQL...", flush=True)
    cols, rows = query(sql)
    print("cols", cols, "nrows", len(rows), flush=True)
    idx = {c: i for i, c in enumerate(cols)}

    stamp = datetime.now().strftime("%Y%m%d")
    out = OUT_DIR / f"max_users_lang_1y_{stamp}.xlsx"
    wb = Workbook()

    # 汇总
    summary = wb.active
    summary.title = "汇总"
    summary.append(["指标", "值"])
    for c in summary[1]:
        c.font = Font(bold=True)
    status_cnt = Counter(r[idx["max_status"]] for r in rows)
    lang_cnt = Counter(r[idx["native_language"]] for r in rows)
    summary.append(["总人数", len(rows)])
    summary.append(["仍在MAX", status_cnt.get("仍在MAX", 0)])
    summary.append(["已退订", status_cnt.get("已退订", 0)])
    summary.append(["口径", "membership=max；含现金+兑换码；不含试用；覆盖期与近365天相交"])
    summary.append([])
    summary.append(["账号母语", "人数", "占比%"])
    for c in summary[7]:
        c.font = Font(bold=True)
    n = len(rows) or 1
    for lang, cnt in lang_cnt.most_common():
        summary.append([lang, cnt, round(cnt * 100.0 / n, 2)])

    # 按状态×语言
    cross = wb.create_sheet("语言x状态")
    cross.append(["账号母语", "仍在MAX", "已退订", "合计"])
    for c in cross[1]:
        c.font = Font(bold=True)
    by_lang_status: dict[str, Counter] = {}
    for r in rows:
        lang = r[idx["native_language"]]
        st = r[idx["max_status"]]
        by_lang_status.setdefault(lang, Counter())[st] += 1
    for lang, c in sorted(by_lang_status.items(), key=lambda x: -sum(x[1].values())):
        a = c.get("仍在MAX", 0)
        b = c.get("已退订", 0)
        cross.append([lang, a, b, a + b])

    # 明细
    detail = wb.create_sheet("明细")
    detail.append([h for _, h in HEADERS])
    for c in detail[1]:
        c.font = Font(bold=True)
    for r in rows:
        detail.append([r[idx[k]] for k, _ in HEADERS])

    wb.save(out)
    print("wrote", out, flush=True)
    print("status", dict(status_cnt), flush=True)
    print("top langs", lang_cnt.most_common(15), flush=True)


if __name__ == "__main__":
    main()
