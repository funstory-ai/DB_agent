# -*- coding: utf-8 -*-
"""T057: T048 口径 + 7d/累计触墙率 → Excel。"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parent))
import export_t018_asof_20260928 as base

ROOT = Path(__file__).resolve().parents[1]
base.SQL_PATH = (
    ROOT
    / "sql"
    / "reports"
    / "query_paywall_ab_longterm_ver_excl_old_hit_20260930.sql"
)
base.OUT_XLSX = (
    ROOT
    / "output"
    / "paywall_ab"
    / "paywall_ab_longterm_ver1326_excl_old_hit_20260930.xlsx"
)


def retitle(path: Path) -> None:
    wb = load_workbook(path)
    bold = Font(bold=True)
    ws = wb["转化总览"]
    ws["A1"] = "付费墙 A/B/C/D · 版本≥1.32.6（语义化）+ 剔老号 · 含触墙率"
    ws["A1"].font = bold
    ws["A2"] = (
        f"asof={base.ASOF}；P040 版本=实验起始 1.32.6（semver）+ P039 剔老号。"
        "触墙=user_need_pro+freeUsageLimitModal；7d=激活后7日内；累计=激活至asof。"
        "对照新付费墙「登录弹窗%」用累计触墙率。"
    )
    if "口径与补算" in wb.sheetnames:
        ws2 = wb["口径与补算"]
        ws2["A1"] = "口径（P040 版本 + P039 剔老号 + 触墙）"
        ws2["A1"].font = bold
        ws2["A20"] = "版本（P040）"
        ws2["B20"] = "语义化 ≥1.32.6（含 1.32.10）"
        ws2["A21"] = "老号（P039）"
        ws2["B21"] = "imt_userinfo.createtime < first_ts - 1 天"
        ws2["A22"] = "触墙"
        ws2["B22"] = (
            "user_need_pro + trigger=freeUsageLimitModal；"
            "7d触墙率=激活后7日内÷激活；触墙率=激活至asof÷激活；"
            "可比新墙 gift 登录弹窗%（观测窗内 show÷激活）"
        )
        ws2["A23"] = "需求"
        ws2["B23"] = "T057；基线 T048"
    wb.save(path)
    print(f"retitled {path}", flush=True)


def main() -> None:
    rows = base.run()
    out = base.write_xlsx(rows)
    retitle(out)
    print("\n=== 7d触墙% / 累计触墙% by week ===", flush=True)
    weeks = [
        "natural_week_0824_0830",
        "natural_week_0831_0906",
        "natural_week_0907_0913",
        "natural_week_0914_0920",
        "natural_week_0921_0927",
        "aggregate_0821_0928",
    ]
    print(
        f"{'slice':28} "
        f"{'A7d':>6} {'Acum':>6} {'B7d':>6} {'Bcum':>6} "
        f"{'C7d':>6} {'Ccum':>6} {'D7d':>6} {'Dcum':>6}",
        flush=True,
    )
    for sl in weeks:
        vals = {}
        for g in "ABCD":
            r = next(x for x in rows if x["slice"] == sl and x["exp_group"] == g)
            vals[g] = (
                float(r["wall_hit_7d_pct"]),
                float(r["wall_hit_to_asof_pct"]),
            )
        print(
            f"{sl:28} "
            f"{vals['A'][0]:5.2f}% {vals['A'][1]:5.2f}% "
            f"{vals['B'][0]:5.2f}% {vals['B'][1]:5.2f}% "
            f"{vals['C'][0]:5.2f}% {vals['C'][1]:5.2f}% "
            f"{vals['D'][0]:5.2f}% {vals['D'][1]:5.2f}%",
            flush=True,
        )
    base.sample_validate(rows)
    # 触墙抽样：A/B 应明显高于 C/D；汇总周人数加总已在 sample_validate
    agg = {r["exp_group"]: r for r in rows if r["slice_type"] == "aggregate"}
    for g in "AB":
        if float(agg[g]["wall_hit_7d_pct"]) <= 0:
            raise SystemExit(f"FAIL {g} 7d触墙应为正")
    for g in "CD":
        if float(agg[g]["wall_hit_7d_pct"]) > 0.5:
            raise SystemExit(f"FAIL {g} 7d触墙应接近0，got {agg[g]['wall_hit_7d_pct']}")
    print("wall sample OK: A/B>0, C/D~0", flush=True)


if __name__ == "__main__":
    main()
