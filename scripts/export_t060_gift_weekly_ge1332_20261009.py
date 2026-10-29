# -*- coding: utf-8 -*-
"""T060 asof=2026-10-09：P034 双口径 · 语义化 version>=1.33.2（新付费墙周刷）。"""
from __future__ import annotations

import sys
from datetime import date as date_cls
from datetime import timedelta, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parent))
from db import query

ROOT = Path(__file__).resolve().parents[1]
SQL_PATH = ROOT / "sql" / "reports" / "query_login_ai_quota_gift_weekly_ge1332_20261009.sql"
OUT_XLSX = (
    ROOT / "output" / "login_ai_quota_gift" / "gift_weekly_p034_dual_ge1332_20261009.xlsx"
)
ASOF = "2026-10-09"
AGG_SLICE = "aggregate_0922_1009"
VERSION_NOTE = "激活行语义化 version>=1.33.2（商店发版起；P040）"
GROUPS = ("A", "B", "C", "I", "J")
SCOPES = (
    ("A_all", "口径A·全量P034"),
    ("B_excl_old_acct", "口径B·剔老号(ACCT<激活-1天)"),
)
SLICE_LABEL = {
    "natural_week_0922_0927": "09/22–27（发版后）",
    "natural_week_0928_1004": "09/28自然周",
    "natural_week_1005_partial": "10/05自然周（本周）",
    AGG_SLICE: "汇总（09-22～10-09）",
}


def _rate(pct_val):
    if pct_val is None:
        return None
    return float(pct_val) / 100.0


def _end_incl(end_excl: str):
    return date_cls.fromisoformat(end_excl) - timedelta(days=1)


def write_overview(ws, agg: dict, title: str, note: str) -> None:
    bold = Font(bold=True)
    pct_fmt = "0.00%"
    ws["A1"] = title
    ws["A1"].font = bold
    ws["A2"] = note
    ws["A3"] = (
        "【分母说明】下表除「激活设备」外，其余列均为百分比，"
        "分母一律=该口径本行「激活设备」UV（不是登录窗 UV，也不是领取 UV）。"
        "例：立即登录%=点了 gift「立即登录」的激活设备/激活；"
        "领取成功%=有 claim 的激活设备/激活（含静默领，故可远高于立即登录%）。"
    )
    funnel_h = [
        "实验组",
        "激活设备(分母)",
        "注册率%",
        "3d checkout%",
        "3d 试用%",
        "3d 付费%",
        "7d 付费%",
        "累计付费%",
        "登录弹窗%",
        "立即登录%",
        "登录关闭%",
        "领取成功%",
        "耗尽墙%",
        "耗尽去付费%",
        "耗尽关闭%",
    ]
    ws["A5"] = "汇总（率=事件UV÷激活设备）"
    ws["A5"].font = bold
    for i, h in enumerate(funnel_h, 1):
        ws.cell(6, i, h).font = bold
    for i, g in enumerate(GROUPS):
        r = 7 + i
        a = agg[g]
        vals = [
            g,
            int(a["new_devices"]),
            _rate(a["register_rate_pct"]),
            _rate(a["checkout_3d_pct"]),
            _rate(a["trial_3d_pct"]),
            _rate(a["paid_3d_pct"]),
            _rate(a["paid_7d_pct"]),
            _rate(a["paid_to_asof_pct"]),
            _rate(a["gift_login_show_pct"]),
            _rate(a["gift_login_click_pct"]),
            _rate(a["gift_login_close_pct"]),
            _rate(a["gift_claim_pct"]),
            _rate(a["gift_exhausted_show_pct"]),
            _rate(a["gift_exhausted_pay_pct"]),
            _rate(a["gift_exhausted_close_pct"]),
        ]
        for col, v in enumerate(vals, 1):
            cell = ws.cell(r, col, v)
            if col >= 3 and v is not None:
                cell.number_format = pct_fmt
    # 列级分母备注
    ws["A13"] = "列分母对照"
    ws["A13"].font = bold
    denom_rows = [
        ("激活设备(分母)", "该口径 P034 激活 UV（A=全量；B=剔老号后）"),
        ("注册率%", "能挂上 deviceinfo.userid 的激活设备 ÷ 激活"),
        ("3d/7d/累计付费%", "对应窗内付费设备 ÷ 激活（付费口径同 T018）"),
        ("登录弹窗%/立即登录%/登录关闭%", "出现对应 gift 登录窗事件的激活设备 ÷ 激活"),
        ("领取成功%", "出现 claim_success 的激活设备 ÷ 激活（含无登录窗的静默领）"),
        ("耗尽墙%/耗尽去付费%/耗尽关闭%", "出现对应耗尽墙事件的激活设备 ÷ 激活"),
        ("注意", "立即登录% 与 领取成功% 不可互为上下界；静默领不经「立即登录」"),
    ]
    for i, (a, b) in enumerate(denom_rows, 14):
        ws.cell(i, 1, a)
        ws.cell(i, 2, b)


def write_weekly(ws, scope_rows: list[dict], title: str) -> None:
    bold = Font(bold=True)
    pct_fmt = "0.00%"
    ws["A1"] = title
    ws["A1"].font = bold
    ws["A2"] = (
        "矩阵「累计付费率」与明细中所有带%列：分母=该同期群×实验组的激活设备 UV。"
        "明细里「登录弹窗show」等无%后缀的是人数（分子 UV），右侧「%」列为÷激活。"
    )

    weekly = [r for r in scope_rows if r["slice_type"] == "weekly"]
    by_slice: dict[str, dict] = {}
    for r in weekly:
        by_slice.setdefault(r["slice"], {"meta": r, "g": {}})
        by_slice[r["slice"]]["g"][r["exp_group"]] = r
    agg = {r["exp_group"]: r for r in scope_rows if r["slice_type"] == "aggregate"}
    slice_order = sorted(by_slice.keys(), key=lambda s: by_slice[s]["meta"]["cohort_no"])
    n_g = len(GROUPS)

    for i, h in enumerate(
        ["同期群", "安装开始", "安装结束", "成熟状态"] + list(GROUPS), 1
    ):
        ws.cell(4, i, h).font = bold
    ws["A3"] = "累计付费率"
    ws["A3"].font = bold

    detail_start = 8 + len(slice_order)
    for si, sk in enumerate(slice_order):
        meta = by_slice[sk]["meta"]
        r = 5 + si
        ws.cell(r, 1, SLICE_LABEL.get(sk, sk))
        ws.cell(r, 2, date_cls.fromisoformat(meta["install_start"]))
        ws.cell(r, 3, _end_incl(meta["install_end_excl"]))
        ws.cell(r, 4, meta["maturity_status"])
        base = detail_start + si * n_g
        for j, _g in enumerate(GROUPS):
            cell = ws.cell(r, 5 + j, f"=M{base + j}")
            cell.number_format = pct_fmt
    mat_agg = 5 + len(slice_order)
    ws.cell(mat_agg, 1, SLICE_LABEL[AGG_SLICE])
    ws.cell(mat_agg, 2, date_cls(2026, 9, 22))
    ws.cell(mat_agg, 3, date_cls(2026, 10, 9))
    ws.cell(mat_agg, 4, "混合成熟度")
    agg_start = detail_start + len(slice_order) * n_g + 2
    for j, _g in enumerate(GROUPS):
        cell = ws.cell(mat_agg, 5 + j, f"=M{agg_start + j}")
        cell.number_format = pct_fmt

    hdr_row = detail_start - 1
    ws.cell(hdr_row - 1, 1, "明细").font = bold
    detail_h = [
        "同期群", "安装开始", "安装结束", "成熟状态", "实验组", "激活设备(分母)",
        "注册率%", "3d checkout%", "3d 试用%", "3d 付费%", "7d 付费%",
        "累计付费人数", "累计付费率%(÷激活)",
        "W1留存%", "W2留存%", "W3留存%", "W4留存%",
        "登录弹窗UV", "登录弹窗%(÷激活)", "立即登录UV", "立即登录%(÷激活)",
        "登录关闭UV", "登录关闭%(÷激活)", "领取UV", "领取%(÷激活)",
        "耗尽墙UV", "耗尽墙%(÷激活)", "耗尽去付费UV", "耗尽付费%(÷激活)",
        "耗尽关闭UV", "耗尽关闭%(÷激活)",
    ]
    for i, h in enumerate(detail_h, 1):
        ws.cell(hdr_row, i, h).font = bold

    def write_detail_row(r, label, start, end_excl, maturity, rec):
        ws.cell(r, 1, label)
        ws.cell(r, 2, date_cls.fromisoformat(start))
        ws.cell(r, 3, _end_incl(end_excl))
        ws.cell(r, 4, maturity)
        ws.cell(r, 5, rec["exp_group"])
        ws.cell(r, 6, int(rec["new_devices"]))
        for col, key in enumerate(
            ["register_rate_pct", "checkout_3d_pct", "trial_3d_pct",
             "paid_3d_pct", "paid_7d_pct"],
            7,
        ):
            cell = ws.cell(r, col, _rate(rec[key]))
            cell.number_format = pct_fmt
        ws.cell(r, 12, int(rec["paid_to_asof"]))
        cell = ws.cell(r, 13, f"=L{r}/F{r}")
        cell.number_format = pct_fmt
        for col, key in enumerate(
            ["ret_w1_pct", "ret_w2_pct", "ret_w3_pct", "ret_w4_pct"], 14
        ):
            v = rec[key]
            if v is None:
                ws.cell(r, col, None)
            else:
                cell = ws.cell(r, col, _rate(v))
                cell.number_format = pct_fmt
        gift_pairs = [
            ("gift_login_show", "gift_login_show_pct"),
            ("gift_login_click", "gift_login_click_pct"),
            ("gift_login_close", "gift_login_close_pct"),
            ("gift_claim", "gift_claim_pct"),
            ("gift_exhausted_show", "gift_exhausted_show_pct"),
            ("gift_exhausted_pay", "gift_exhausted_pay_pct"),
            ("gift_exhausted_close", "gift_exhausted_close_pct"),
        ]
        col = 18
        for cnt_k, pct_k in gift_pairs:
            ws.cell(r, col, int(rec[cnt_k] or 0))
            cell = ws.cell(r, col + 1, _rate(rec[pct_k]))
            cell.number_format = pct_fmt
            col += 2

    ri = detail_start
    for sk in slice_order:
        meta = by_slice[sk]["meta"]
        for g in GROUPS:
            write_detail_row(
                ri,
                SLICE_LABEL.get(sk, sk),
                meta["install_start"],
                meta["install_end_excl"],
                meta["maturity_status"],
                by_slice[sk]["g"][g],
            )
            ri += 1
    ri = agg_start
    for g in GROUPS:
        write_detail_row(
            ri,
            SLICE_LABEL[AGG_SLICE],
            "2026-09-22",
            "2026-10-10",
            "混合成熟度",
            agg[g],
        )
        ri += 1


def sample_validate() -> bool:
    """短窗（北京 10-09 一天）抽 20：验 P039 1 天缓冲；不重扫全 cohort。"""
    print("\n【抽样复核】老号 1 天缓冲（短窗 10-09）...", flush=True)
    sql_old = r"""
    WITH act AS (
      SELECT device_id, created_ts AS act_ts, TRY_CAST(ex_int_arg1 AS BIGINT) AS act_uid
      FROM (
        SELECT device_id, created_ts, ex_int_arg1,
          row_number() OVER (PARTITION BY device_id ORDER BY created_ts) rn
        FROM immdata.user_behaviour
        WHERE created_ts >= '2026-10-08 16:00:00' AND created_ts < '2026-10-09 16:00:00'
          AND event_name = 'translate_page' AND is_new_user_today = '1'
          AND lower(trim(ab_tag)) = 'j'
          AND (
            TRY_CAST(split_part(version, '.', 1) AS INT) > 1
            OR (
              TRY_CAST(split_part(version, '.', 1) AS INT) = 1
              AND (
                TRY_CAST(split_part(version, '.', 2) AS INT) > 33
                OR (
                  TRY_CAST(split_part(version, '.', 2) AS INT) = 33
                  AND TRY_CAST(split_part(version, '.', 3) AS INT) >= 2
                )
              )
            )
          )
      ) t WHERE rn = 1
    ),
    di AS (
      SELECT a.device_id, a.act_ts, d.userid AS uid
      FROM act a
      INNER JOIN immdata.imt_deviceinfo d
        ON d.deviceid = a.device_id AND d.userid IS NOT NULL AND COALESCE(d.isdeleted,0)=0
    ),
    flagged AS (
      SELECT di.device_id, di.act_ts, di.uid, u.createtime AS acct_ct
      FROM di
      INNER JOIN immdata.imt_userinfo u ON u.id = di.uid AND COALESCE(u.isdeleted,0)=0
      WHERE u.createtime < date_add(day, -1, di.act_ts)
      LIMIT 20
    )
    SELECT
      count(*) AS n,
      sum(IF(acct_ct < date_add(day, -1, act_ts), 1, 0)) AS n_ok_rule,
      sum(IF(acct_ct >= date_add(day, -1, act_ts), 1, 0)) AS n_bad
    FROM flagged
    """
    sql_kept = r"""
    WITH act AS (
      SELECT device_id, created_ts AS act_ts
      FROM (
        SELECT device_id, created_ts,
          row_number() OVER (PARTITION BY device_id ORDER BY created_ts) rn
        FROM immdata.user_behaviour
        WHERE created_ts >= '2026-10-08 16:00:00' AND created_ts < '2026-10-09 16:00:00'
          AND event_name = 'translate_page' AND is_new_user_today = '1'
          AND lower(trim(ab_tag)) = 'j'
          AND (
            TRY_CAST(split_part(version, '.', 1) AS INT) > 1
            OR (
              TRY_CAST(split_part(version, '.', 1) AS INT) = 1
              AND (
                TRY_CAST(split_part(version, '.', 2) AS INT) > 33
                OR (
                  TRY_CAST(split_part(version, '.', 2) AS INT) = 33
                  AND TRY_CAST(split_part(version, '.', 3) AS INT) >= 2
                )
              )
            )
          )
      ) t WHERE rn = 1
    ),
    di AS (
      SELECT a.device_id, a.act_ts, d.userid AS uid
      FROM act a
      INNER JOIN immdata.imt_deviceinfo d
        ON d.deviceid = a.device_id AND d.userid IS NOT NULL AND COALESCE(d.isdeleted,0)=0
    ),
    kept AS (
      SELECT di.device_id, di.act_ts, di.uid, u.createtime AS acct_ct
      FROM di
      INNER JOIN immdata.imt_userinfo u ON u.id = di.uid AND COALESCE(u.isdeleted,0)=0
      WHERE u.createtime >= date_add(day, -1, di.act_ts)
      LIMIT 20
    )
    SELECT
      count(*) AS n,
      sum(IF(acct_ct >= date_add(day, -1, act_ts), 1, 0)) AS n_ok_rule,
      sum(IF(acct_ct < date_add(day, -1, act_ts), 1, 0)) AS n_bad
    FROM kept
    """
    ok = True
    for kind, sql in (("old_flag", sql_old), ("kept_reg", sql_kept)):
        cols, rows = query(sql)
        d = dict(zip(cols, rows[0]))
        d["kind"] = kind
        print(d, flush=True)
        if int(d["n"] or 0) == 0 or int(d["n_bad"] or 0) > 0:
            ok = False
    return ok


def main() -> None:
    sql = SQL_PATH.read_text(encoding="utf-8")
    print("querying dual-scope weekly...", flush=True)
    cols, rows = query(sql)
    out = [dict(zip(cols, r)) for r in rows]

    by_scope: dict[str, list] = {}
    for r in out:
        by_scope.setdefault(r["metric_scope"], []).append(r)

    for scope, label in SCOPES:
        print(f"\n=== {scope} {label} ===", flush=True)
        for r in by_scope[scope]:
            if r["slice_type"] != "aggregate":
                continue
            print(
                r["exp_group"],
                "n=", r["new_devices"],
                "reg=", r["register_rate_pct"],
                "paid=", r["paid_to_asof_pct"],
                "claim=", r.get("gift_claim"),
                flush=True,
            )

    wb = Workbook()
    bold = Font(bold=True)

    # first overview
    scope0, label0 = SCOPES[0]
    agg0 = {
        r["exp_group"]: r
        for r in by_scope[scope0]
        if r["slice_type"] == "aggregate"
    }
    ws = wb.active
    ws.title = "总览_A全量"
    write_overview(
        ws,
        agg0,
        f"登录赠送额度（T060）· {label0} · {VERSION_NOTE}",
        f"口径A分母=P034激活且{VERSION_NOTE}；I=i对照，J=j实验；激活日起09-22；asof={ASOF}。"
        "表内率列均为百分比格式，分母见 A3/列分母对照。",
    )

    scope1, label1 = SCOPES[1]
    agg1 = {
        r["exp_group"]: r
        for r in by_scope[scope1]
        if r["slice_type"] == "aggregate"
    }
    ws_b = wb.create_sheet("总览_B剔老号")
    write_overview(
        ws_b,
        agg1,
        f"登录赠送额度（T060）· {label1} · {VERSION_NOTE}",
        f"口径B分母=P034激活×{VERSION_NOTE}且剔老号（createtime<激活-1天）；asof={ASOF}。"
        "表内率列均为百分比格式，分母见 A3/列分母对照。",
    )

    write_weekly(
        wb.create_sheet("周_A全量"),
        by_scope[scope0],
        f"周同期群 · {label0}",
    )
    write_weekly(
        wb.create_sheet("周_B剔老号"),
        by_scope[scope1],
        f"周同期群 · {label1}",
    )

    ws2 = wb.create_sheet("口径")
    ws2["A1"] = "口径"
    ws2["A1"].font = bold
    notes = [
        ("版本", "激活首条 translate_page 语义化 version>=1.33.2（P040）"),
        ("口径A分母", "P034 激活设备 UV（与 T018 对齐）× 上列版本"),
        ("口径B分母", "A 剔除老号：userid.createtime < first_ts−1天（deviceinfo∪激活行∪claim）"),
        ("百分比约定", "总览/周表明细中带%的列一律 Excel 百分比格式；人数列不带%"),
        ("统一率分母", "注册/付费/gift登录/领取/耗尽 各率分母=该口径该切片的激活设备，不是上一漏斗步"),
        ("立即登录% vs 领取%", "立即登录只计 gift 登录窗主按钮；领取含别处登录后静默领，故领取%可≫立即登录%"),
        ("1天缓冲", "同会话先注册再译、deviceinfo 入表延迟不误伤；不用 deviceinfo 时间判老号"),
        ("分组", "A=a/b/c；B=d/e/f；C=g/h；I=i；J=j（loginAiQuotaGift）"),
        ("起点", "激活北京日 ≥ 2026-09-22；有效包自 1.33.2 起；1.33.3 放量更大"),
        ("注册", "LEFT JOIN imt_deviceinfo.userid；禁止 INNER（P031）"),
        ("付费", "T018 v3：finished 现金/onetime/code；剔 iospay_renew 与锚点前 finished"),
        ("额度机制", "奖励中心 task.taken 跟账号；本机 storage 另有设备锁防换号再领"),
    ]
    for i, (a, b) in enumerate(notes, 3):
        ws2.cell(i, 1, a)
        ws2.cell(i, 2, b)

    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)
    print(f"\nwrote {OUT_XLSX}", flush=True)

    ok = True
    agg0 = {
        r["exp_group"]: r
        for r in by_scope["A_all"]
        if r["slice_type"] == "aggregate"
    }
    agg1 = {
        r["exp_group"]: r
        for r in by_scope["B_excl_old_acct"]
        if r["slice_type"] == "aggregate"
    }
    for scope, _label in SCOPES:
        scope_rows = by_scope[scope]
        weekly = [r for r in scope_rows if r["slice_type"] == "weekly"]
        agg = {
            r["exp_group"]: r
            for r in scope_rows
            if r["slice_type"] == "aggregate"
        }
        for g in GROUPS:
            s = sum(int(r["new_devices"]) for r in weekly if r["exp_group"] == g)
            a = int(agg[g]["new_devices"])
            if s != a:
                print(f"FAIL {scope} sum {g} weekly={s} agg={a}")
                ok = False
        if int(agg["J"]["gift_login_show"] or 0) <= 0:
            print(f"FAIL {scope} J should have login_show")
            ok = False
    a_n = int(agg0["J"]["new_devices"])
    b_n = int(agg1["J"]["new_devices"])
    if b_n >= a_n:
        print(f"FAIL B should shrink vs A for J: A={a_n} B={b_n}")
        ok = False
    else:
        print(
            f"J activate A={a_n} B={b_n} excluded={a_n - b_n} "
            f"({(a_n - b_n) * 100.0 / a_n:.1f}%)",
            flush=True,
        )
    for g in GROUPS:
        if int(agg1[g]["new_devices"]) > int(agg0[g]["new_devices"]):
            print(f"FAIL B>A for {g}")
            ok = False

    # last_closed=09-28：09/22 周应有 W1；09/28、10/05 周 W1 应空
    for scope, _label in SCOPES:
        for r in by_scope[scope]:
            if r["slice"] == "natural_week_0922_0927" and r["ret_w1_pct"] is None:
                print(f"FAIL {scope} 09/22 should have W1", r["exp_group"])
                ok = False
            if r["slice"] in (
                "natural_week_0928_1004",
                "natural_week_1005_partial",
            ) and r["ret_w1_pct"] is not None:
                print(f"FAIL {scope} {r['slice']} W1 should be empty", r["exp_group"])
                ok = False

    sample_ok = sample_validate()
    ok = ok and sample_ok
    print(f"【抽样复核】结论={'通过' if ok else '不通过'}", flush=True)
    if not ok:
        sys.exit(1)


if __name__ == "__main__":
    main()
