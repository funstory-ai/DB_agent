# -*- coding: utf-8 -*-
"""T065 新版 50%：激活≥2026-10-10，asof=10-12；无墙组(a-e) / 有墙组(f-j)。"""
from __future__ import annotations

import sys
from datetime import date as date_cls
from datetime import timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parent))
from db import query

ROOT = Path(__file__).resolve().parents[1]
SQL_PATH = (
    ROOT
    / "sql"
    / "reports"
    / "query_login_ai_quota_gift_weekly_ge1332_50pct_from1010_20261012.sql"
)
OUT_XLSX = (
    ROOT
    / "output"
    / "login_ai_quota_gift"
    / "gift_weekly_p034_dual_ge1332_50pct_from1010_20261012.xlsx"
)
ASOF = "2026-10-12"
AGG_SLICE = "aggregate_1010_1012"
VERSION_NOTE = "激活行语义化 version>=1.33.2（P040）"
GROUPS = ("无墙组(a-e)", "有墙组(f-j)")
SCOPES = (
    ("A_all", "口径A·全量P034"),
    ("B_excl_old_acct", "口径B·剔老号(ACCT<激活-1天)"),
)
SLICE_LABEL = {
    "natural_week_1010_1011": "10/10–11（扩量后）",
    "natural_week_1012_partial": "10/12（本周1天）",
    AGG_SLICE: "汇总（10-10～10-12·50%）",
}
AGG_START = "2026-10-10"
AGG_END_EXCL = "2026-10-13"
AGG_END_INCL = date_cls(2026, 10, 12)


def _rate(pct_val):
    if pct_val is None:
        return None
    return float(pct_val) / 100.0


def _end_incl(end_excl: str):
    return date_cls.fromisoformat(end_excl) - timedelta(days=1)


def write_overview(ws, agg: dict, title: str, note: str) -> None:
    bold = Font(bold=True)
    pct_fmt = "0.00%"
    ws["A1"] = title
    ws["A1"].font = bold
    ws["A2"] = note
    funnel_h = [
        "实验组",
        "激活设备(分母)",
        "注册率%",
        "3d checkout%",
        "3d 试用%",
        "3d 付费%",
        "7d 付费%",
        "累计付费人数(总付费)",
        "累计付费%",
        "登录弹窗%",
        "立即登录%",
        "登录关闭%",
        "领取成功%",
        "耗尽墙%",
        "耗尽去付费人数",
        "耗尽去付费%",
        "去付费后付费人数",
        "去付费→付费率%(÷去付费)",
        "耗尽关闭%",
    ]
    ws["A5"] = "汇总（无墙组=a–e；有墙组=f–j；对齐 abTagNumber>4 约 50%）"
    ws["A5"].font = bold
    for i, h in enumerate(funnel_h, 1):
        ws.cell(6, i, h).font = bold
    for i, g in enumerate(GROUPS):
        r = 7 + i
        a = agg[g]
        vals = [
            g,
            int(a["new_devices"]),
            _rate(a["register_rate_pct"]),
            _rate(a["checkout_3d_pct"]),
            _rate(a["trial_3d_pct"]),
            _rate(a["paid_3d_pct"]),
            _rate(a["paid_7d_pct"]),
            int(a["paid_to_asof"] or 0),
            _rate(a["paid_to_asof_pct"]),
            _rate(a["gift_login_show_pct"]),
            _rate(a["gift_login_click_pct"]),
            _rate(a["gift_login_close_pct"]),
            _rate(a["gift_claim_pct"]),
            _rate(a["gift_exhausted_show_pct"]),
            int(a["gift_exhausted_pay"] or 0),
            _rate(a["gift_exhausted_pay_pct"]),
            int(a["gift_exhausted_pay_then_paid"] or 0),
            _rate(a["gift_exhausted_pay_then_paid_pct"]),
            _rate(a["gift_exhausted_close_pct"]),
        ]
        for col, v in enumerate(vals, 1):
            cell = ws.cell(r, col, v)
            if col in (3, 4, 5, 6, 7, 9, 10, 11, 12, 13, 14, 16, 18, 19) and v is not None:
                cell.number_format = pct_fmt


def write_weekly(ws, scope_rows: list[dict], title: str) -> None:
    bold = Font(bold=True)
    pct_fmt = "0.00%"
    ws["A1"] = title
    ws["A1"].font = bold

    weekly = [r for r in scope_rows if r["slice_type"] == "weekly"]
    by_slice: dict[str, dict] = {}
    for r in weekly:
        by_slice.setdefault(r["slice"], {"meta": r, "g": {}})
        by_slice[r["slice"]]["g"][r["exp_group"]] = r
    agg = {r["exp_group"]: r for r in scope_rows if r["slice_type"] == "aggregate"}
    slice_order = sorted(by_slice.keys(), key=lambda s: by_slice[s]["meta"]["cohort_no"])
    n_g = len(GROUPS)

    for i, h in enumerate(
        ["同期群", "安装开始", "安装结束", "成熟状态"] + list(GROUPS), 1
    ):
        ws.cell(4, i, h).font = bold
    ws["A3"] = "累计付费率"
    ws["A3"].font = bold

    detail_start = 8 + len(slice_order)
    for si, sk in enumerate(slice_order):
        meta = by_slice[sk]["meta"]
        r = 5 + si
        ws.cell(r, 1, SLICE_LABEL.get(sk, sk))
        ws.cell(r, 2, date_cls.fromisoformat(meta["install_start"]))
        ws.cell(r, 3, _end_incl(meta["install_end_excl"]))
        ws.cell(r, 4, meta["maturity_status"])
        base = detail_start + si * n_g
        for j, _g in enumerate(GROUPS):
            cell = ws.cell(r, 5 + j, f"=M{base + j}")
            cell.number_format = pct_fmt
    mat_agg = 5 + len(slice_order)
    ws.cell(mat_agg, 1, SLICE_LABEL[AGG_SLICE])
    ws.cell(mat_agg, 2, date_cls.fromisoformat(AGG_START))
    ws.cell(mat_agg, 3, AGG_END_INCL)
    ws.cell(mat_agg, 4, "混合成熟度·50%")
    agg_start = detail_start + len(slice_order) * n_g + 2
    for j, _g in enumerate(GROUPS):
        cell = ws.cell(mat_agg, 5 + j, f"=M{agg_start + j}")
        cell.number_format = pct_fmt

    hdr_row = detail_start - 1
    ws.cell(hdr_row - 1, 1, "明细").font = bold
    detail_h = [
        "同期群", "安装开始", "安装结束", "成熟状态", "实验组", "激活设备(分母)",
        "注册率%", "3d checkout%", "3d 试用%", "3d 付费%", "7d 付费%",
        "累计付费人数", "累计付费率%(÷激活)",
        "W1留存%", "W2留存%", "W3留存%", "W4留存%",
        "登录弹窗UV", "登录弹窗%(÷激活)", "立即登录UV", "立即登录%(÷激活)",
        "登录关闭UV", "登录关闭%(÷激活)", "领取UV", "领取%(÷激活)",
        "耗尽墙UV", "耗尽墙%(÷激活)", "耗尽去付费UV", "耗尽去付费%(÷激活)",
        "去付费后付费人数", "去付费→付费率%(÷去付费)",
        "耗尽关闭UV", "耗尽关闭%(÷激活)",
    ]
    for i, h in enumerate(detail_h, 1):
        ws.cell(hdr_row, i, h).font = bold

    def write_detail_row(r, label, start, end_excl, maturity, rec):
        ws.cell(r, 1, label)
        ws.cell(r, 2, date_cls.fromisoformat(start))
        ws.cell(r, 3, _end_incl(end_excl))
        ws.cell(r, 4, maturity)
        ws.cell(r, 5, rec["exp_group"])
        ws.cell(r, 6, int(rec["new_devices"]))
        for col, key in enumerate(
            ["register_rate_pct", "checkout_3d_pct", "trial_3d_pct",
             "paid_3d_pct", "paid_7d_pct"],
            7,
        ):
            cell = ws.cell(r, col, _rate(rec[key]))
            cell.number_format = pct_fmt
        ws.cell(r, 12, int(rec["paid_to_asof"]))
        cell = ws.cell(r, 13, f"=L{r}/F{r}")
        cell.number_format = pct_fmt
        for col, key in enumerate(
            ["ret_w1_pct", "ret_w2_pct", "ret_w3_pct", "ret_w4_pct"], 14
        ):
            v = rec[key]
            if v is None:
                ws.cell(r, col, None)
            else:
                cell = ws.cell(r, col, _rate(v))
                cell.number_format = pct_fmt
        gift_pairs = [
            ("gift_login_show", "gift_login_show_pct"),
            ("gift_login_click", "gift_login_click_pct"),
            ("gift_login_close", "gift_login_close_pct"),
            ("gift_claim", "gift_claim_pct"),
            ("gift_exhausted_show", "gift_exhausted_show_pct"),
            ("gift_exhausted_pay", "gift_exhausted_pay_pct"),
        ]
        col = 18
        for cnt_k, pct_k in gift_pairs:
            ws.cell(r, col, int(rec[cnt_k] or 0))
            cell = ws.cell(r, col + 1, _rate(rec[pct_k]))
            cell.number_format = pct_fmt
            col += 2
        ws.cell(r, col, int(rec["gift_exhausted_pay_then_paid"] or 0))
        cell = ws.cell(r, col + 1, _rate(rec["gift_exhausted_pay_then_paid_pct"]))
        if rec["gift_exhausted_pay_then_paid_pct"] is not None:
            cell.number_format = pct_fmt
        col += 2
        ws.cell(r, col, int(rec["gift_exhausted_close"] or 0))
        cell = ws.cell(r, col + 1, _rate(rec["gift_exhausted_close_pct"]))
        cell.number_format = pct_fmt

    ri = detail_start
    for sk in slice_order:
        meta = by_slice[sk]["meta"]
        for g in GROUPS:
            write_detail_row(
                ri,
                SLICE_LABEL.get(sk, sk),
                meta["install_start"],
                meta["install_end_excl"],
                meta["maturity_status"],
                by_slice[sk]["g"][g],
            )
            ri += 1
    ri = agg_start
    for g in GROUPS:
        write_detail_row(
            ri,
            SLICE_LABEL[AGG_SLICE],
            AGG_START,
            AGG_END_EXCL,
            "混合成熟度·50%",
            agg[g],
        )
        ri += 1


def sample_validate() -> bool:
    """短窗 10-11：有墙组应有墙，无墙组近 0。"""
    print("\n【抽样复核】50% 扩量后 无墙组 vs 有墙组…", flush=True)
    sql = r"""
    SELECT
      count(DISTINCT CASE WHEN lower(trim(ab_tag)) IN ('a','b','c','d','e')
        THEN device_id END) AS ctrl_show,
      count(DISTINCT CASE WHEN lower(trim(ab_tag)) IN ('f','g','h','i','j')
        THEN device_id END) AS treat_show
    FROM immdata.user_behaviour
    WHERE created_ts >= '2026-10-11 16:00:00' AND created_ts < '2026-10-12 16:00:00'
      AND event_name = 'login_ai_quota_gift_login_modal_show'
      AND ab_tag IS NOT NULL AND length(trim(ab_tag)) = 1
    """
    cols, rows = query(sql)
    d = dict(zip(cols, rows[0]))
    print(d, flush=True)
    ok = int(d["treat_show"] or 0) > 1000
    if int(d["ctrl_show"] or 0) > int(d["treat_show"] or 0) * 0.1:
        print("FAIL 无墙组 gift show unexpectedly high")
        ok = False
    return ok


def main() -> None:
    sql = SQL_PATH.read_text(encoding="utf-8")
    print("querying 50% from 10-10...", flush=True)
    cols, rows = query(sql)
    out = [dict(zip(cols, r)) for r in rows]
    by_scope: dict[str, list] = {}
    for r in out:
        by_scope.setdefault(r["metric_scope"], []).append(r)

    for scope, label in SCOPES:
        print(f"\n=== {scope} {label} ===", flush=True)
        for r in by_scope[scope]:
            if r["slice_type"] != "aggregate":
                continue
            print(
                r["exp_group"],
                "n=", r["new_devices"],
                "login%=", r["gift_login_show_pct"],
                "exh_pay=", r.get("gift_exhausted_pay"),
                "exh→paid=", r.get("gift_exhausted_pay_then_paid"),
                flush=True,
            )

    wb = Workbook()
    bold = Font(bold=True)
    scope0, label0 = SCOPES[0]
    agg0 = {r["exp_group"]: r for r in by_scope[scope0] if r["slice_type"] == "aggregate"}
    ws = wb.active
    ws.title = "总览_A全量"
    write_overview(
        ws,
        agg0,
        f"登录赠送·新版50%（T065）· {label0}",
        f"激活日≥2026-10-10；asof={ASOF}；{VERSION_NOTE}；"
        "无墙组=a–e / 有墙组=f–j（abTagNumber>4）。与旧版仅 j 隔离。",
    )
    scope1, label1 = SCOPES[1]
    agg1 = {r["exp_group"]: r for r in by_scope[scope1] if r["slice_type"] == "aggregate"}
    write_overview(
        wb.create_sheet("总览_B剔老号"),
        agg1,
        f"登录赠送·新版50%（T065）· {label1}",
        f"口径B剔老号；激活≥10-10；asof={ASOF}。",
    )
    write_weekly(wb.create_sheet("周_A全量"), by_scope[scope0], f"周同期群 · {label0}")
    write_weekly(wb.create_sheet("周_B剔老号"), by_scope[scope1], f"周同期群 · {label1}")

    ws2 = wb.create_sheet("口径")
    ws2["A1"] = "口径"
    ws2["A1"].font = bold
    notes = [
        ("阶段", "新版 50%：09639bb4c 起 abTagNumber>4（f–j）"),
        ("旧版", "见 T064 / gift_*_jonly_asof1009_*（仅 j，asof=10-09）"),
        ("分组", "无墙组=a/b/c/d/e；有墙组=f/g/h/i/j"),
        ("版本", VERSION_NOTE),
        ("口径B", "P039 剔老号"),
        ("去付费→付费", "P042"),
        ("需求", "T065"),
    ]
    for i, (a, b) in enumerate(notes, 3):
        ws2.cell(i, 1, a)
        ws2.cell(i, 2, b)

    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)
    print(f"\nwrote {OUT_XLSX}", flush=True)

    ok = True
    for scope, _label in SCOPES:
        scope_rows = by_scope[scope]
        weekly = [r for r in scope_rows if r["slice_type"] == "weekly"]
        agg = {r["exp_group"]: r for r in scope_rows if r["slice_type"] == "aggregate"}
        for g in GROUPS:
            s = sum(int(r["new_devices"]) for r in weekly if r["exp_group"] == g)
            a = int(agg[g]["new_devices"])
            if s != a:
                print(f"FAIL {scope} sum {g} weekly={s} agg={a}")
                ok = False
        g_ctrl, g_treat = GROUPS
        ctrl_pct = float(agg[g_ctrl]["gift_login_show_pct"] or 0)
        treat_pct = float(agg[g_treat]["gift_login_show_pct"] or 0)
        if treat_pct < 20:
            print(f"FAIL {scope} 有墙组 login%={treat_pct} too low for 50%")
            ok = False
        if ctrl_pct > 5:
            print(f"FAIL {scope} 无墙组 login%={ctrl_pct} should be ~0")
            ok = False
        # 1:1 量级
        cn = int(agg[g_ctrl]["new_devices"])
        tn = int(agg[g_treat]["new_devices"])
        if cn + tn > 0:
            share = tn / (cn + tn)
            if share < 0.35 or share > 0.65:
                print(f"WARN {scope} 有墙组 share={share:.2%} (expect ~50%)")
        t = agg[g_treat]
        exh = int(t["gift_exhausted_pay"] or 0)
        then_paid = int(t["gift_exhausted_pay_then_paid"] or 0)
        total_paid = int(t["paid_to_asof"] or 0)
        if then_paid > exh or then_paid > total_paid:
            print(f"FAIL {scope} 有墙组 then_paid inconsistent")
            ok = False
        print(
            f"{scope} 无墙组 n={cn} login%={ctrl_pct} | "
            f"有墙组 n={tn} login%={treat_pct} 去付费={exh}→付费={then_paid}",
            flush=True,
        )
        # 两切片均未满 W1
        for r in scope_rows:
            if r["slice_type"] == "weekly" and r["ret_w1_pct"] is not None:
                print(f"FAIL {scope} {r['slice']} W1 should be empty (immature)")
                ok = False

    a_n = int(agg0[GROUPS[1]]["new_devices"])
    b_n = int(agg1[GROUPS[1]]["new_devices"])
    if b_n >= a_n:
        print("FAIL B should shrink vs A for 有墙组")
        ok = False

    ok = ok and sample_validate()
    print(f"【抽样复核】结论={'通过' if ok else '不通过'}", flush=True)
    if not ok:
        sys.exit(1)


if __name__ == "__main__":
    main()
