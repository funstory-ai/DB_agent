# -*- coding: utf-8 -*-
"""T070 付费用户付款前行为。付费含兑换码，不含试用。截止北京 2026-10-16 01:05。"""
from __future__ import annotations

import sys
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

sys.path.insert(0, str(Path(__file__).resolve().parent))
import export_t070_narrow_funnel_20261014 as m

# 订单副本最新成交在北京 2026-10-16 01:05（日更水位）
m.OBS_END = datetime(2026, 10, 15, 17, 6, 0)
INSTALL_END = "2026-10-17"
OUT = (
    m.ROOT
    / "output"
    / "login_ai_quota_gift"
    / "gift_paid_behaviors_20261016.xlsx"
)
EVENTS = (
    "go_pricing_page",
    "pricing_purchase_btn_click",
    "purchase_success",
    "user_need_pro",
    "change_translate_service",
    "translate_video_subtitle",
    "use_ai_subtitle",
    "open_pdf_page",
    "login_ai_quota_gift_claim_success",
    "login_ai_quota_gift_login_modal_show",
    "login_ai_quota_gift_exhausted_modal_show",
)


def q(sql):
    return m.query(sql)[1]


def beijing(ts):
    return (ts + timedelta(hours=8)).strftime("%m-%d %H:%M")


def label(name, trigger):
    t = trigger or ""
    mapping = {
        "login_ai_quota_gift_claim_success": "领了赠送",
        "login_ai_quota_gift_login_modal_show": "登录墙",
        "login_ai_quota_gift_exhausted_modal_show": "耗尽墙",
        "translate_video_subtitle": "看视频字幕",
        "use_ai_subtitle": "用AI字幕",
        "open_pdf_page": "打开PDF",
        "change_translate_service": "切换翻译服务",
    }
    if name in mapping:
        return mapping[name]
    if name == "user_need_pro":
        return f"会员墙[{t}]" if t else "会员墙"
    if name == "go_pricing_page":
        return f"去定价[{t}]" if t else "去定价"
    if name == "pricing_purchase_btn_click":
        return f"点购买[{t}]" if t else "点购买"
    if name == "purchase_success":
        return f"购买成功[{t}]" if t else "购买成功"
    return name


def specific_bucket(text):
    t = (text or "").lower()
    if any(k in t for k in ("ai_subtitle", "subtitle", "no_subtitle")):
        return "视频或AI字幕"
    if any(k in t for k in ("upgrade_to_pro", "upgrade_to_max", "upgrade_pro", "upgrade_max", "login_use_pro", "login_pro")):
        return "切到 Pro/Max"
    if any(k in t for k in ("pdf", "babel", "scanned", "journal")):
        return "PDF/文档"
    if any(k in t for k in ("image", "manga")):
        return "漫画/图片"
    if "service_error" in t:
        return "翻译服务报错"
    if "ai_context" in t:
        return "AI上下文"
    if "onboarding" in t:
        return "新手引导"
    return None


def classify(events, pay_ts):
    ev = [e for e in events if e[0] <= pay_ts + timedelta(minutes=1)]
    candidates = []
    for ts, name, trigger in ev:
        if name not in ("purchase_success", "pricing_purchase_btn_click", "go_pricing_page", "user_need_pro"):
            continue
        hit = specific_bucket(trigger)
        if hit:
            candidates.append((ts, hit, f"最靠近付款 {label(name, trigger)}"))
    if candidates:
        _ts, hit, why = max(candidates, key=lambda x: x[0])
        return hit, why
    names = {e[1] for e in ev}
    if "use_ai_subtitle" in names or "translate_video_subtitle" in names:
        return "视频或AI字幕", "看过视频或用过AI字幕"
    if "open_pdf_page" in names:
        return "PDF/文档", "打开过PDF"
    if "change_translate_service" in names:
        return "切到 Pro/Max", "切换了翻译服务"
    if any("onboarding" in (e[2] or "") for e in ev):
        return "新手引导", "新手引导里的定价"
    if any(
        (e[2] or "") in ("profile", "popup_user_upgrade", "float_ball_popup_user_upgrade", "extension_settings_general")
        or e[1] == "go_pricing_page"
        for e in ev
        if e[1] in ("purchase_success", "pricing_purchase_btn_click", "go_pricing_page")
    ):
        return "个人页或弹窗", "从个人页或弹窗去定价，没有明确功能墙"
    if "login_ai_quota_gift_claim_success" in names:
        return "只看到领了赠送", "只有领取"
    return "没有对上功能事件", "窗口内没有上述事件"


def collapse(labels):
    out = []
    for item in labels:
        if not out or out[-1][0] != item:
            out.append([item, 1])
        else:
            out[-1][1] += 1
    parts = []
    for item, n in out:
        parts.append(item if n == 1 else f"{item}×{n}")
    return " → ".join(parts)


def main():
    cohort = m.load_cohort(INSTALL_END)
    gift = m.load_gift(cohort)
    m.load_device_users(cohort)
    order_rows = q(f"""
    SELECT o.userid, o.createtime, o.amountpaid, lower(COALESCE(o.ordertype, '')),
      CASE
        WHEN CAST(o.subscriptiongoodssnap AS STRING) LIKE '%"SubscriptionType":"onetime%'
          OR lower(COALESCE(o.ordertype, '')) = 'code'
        THEN 1 ELSE 0
      END
    FROM immdata.imt_subscriptiongoodsorder o
    WHERE COALESCE(o.isdeleted, 0) = 0
      AND o.orderstatus = 'finished'
      AND o.userid IS NOT NULL
      AND lower(COALESCE(o.ordertype, '')) <> 'iospay_renew'
      AND o.createtime >= '{m.ORDER_START:%Y-%m-%d %H:%M:%S}'
      AND o.createtime < '{m.OBS_END:%Y-%m-%d %H:%M:%S}'
    """)
    orders = defaultdict(list)
    for uid, ts, amount, otype, is_code in order_rows:
        amount_num = 0 if amount is None else int(amount)
        is_paid = amount_num > 0 or int(is_code or 0) == 1
        if is_paid:
            orders[int(uid)].append((ts, amount, otype, int(is_code or 0), amount_num > 0))

    uids = set()
    for c in cohort.values():
        if c["userid"]:
            uids.add(int(c["userid"]))
        if c["act_uid"]:
            uids.add(int(c["act_uid"]))
    for evs in gift.values():
        for _ev, _ts, uid in evs:
            if uid:
                uids.add(int(uid))
    userinfo = m.load_userinfo(uids)

    devices_by_uid = defaultdict(list)
    for device_id, c in cohort.items():
        if c["grp"] != m.GROUPS[1]:
            continue
        if c["install_day"] < "2026-10-10" or c["install_day"] >= INSTALL_END:
            continue
        linked = set()
        for u in (c["userid"], c["act_uid"]):
            if u:
                linked.add(int(u))
        for _ev, _ts, uid in gift.get(device_id, []):
            if uid:
                linked.add(int(uid))
        old = False
        for u in linked:
            ct = userinfo.get(u)
            if ct is not None and ct < c["first_ts"] - timedelta(days=1):
                old = True
                break
        if old:
            continue
        for u in (c["userid"], c["act_uid"]):
            if u:
                devices_by_uid[int(u)].append(device_id)

    people = {}
    for uid, devs in devices_by_uid.items():
        firsts = [cohort[d]["first_ts"] for d in devs]
        anchor = min(firsts)
        paid = [o for o in orders.get(uid, []) if anchor <= o[0] < m.OBS_END]
        if not paid:
            continue
        paid.sort(key=lambda o: o[0])
        pay_ts, amount, otype, is_code, is_cash = paid[0]
        exh = []
        for d in devs:
            exh.extend(ts for ev, ts, _u in gift.get(d, []) if ev == m.EXH and ts <= pay_ts)
        earliest = min(devs, key=lambda d: cohort[d]["first_ts"])
        people[uid] = {
            "ab": cohort[earliest]["ab"],
            "first_ts": anchor,
            "pay_ts": pay_ts,
            "amount": amount,
            "otype": otype,
            "kind": "应用内付款" if is_cash else "兑换码",
            "wall": bool(exh),
            "devs": devs,
            "claim": any(ev == m.CLAIM for d in devs for ev, _ts, _u in gift.get(d, [])),
            "login": any(ev == m.SHOW for d in devs for ev, _ts, _u in gift.get(d, [])),
        }
    print("paid accounts", len(people), "wall", sum(1 for p in people.values() if p["wall"]), flush=True)

    dev_ids = sorted({d for p in people.values() for d in p["devs"]})
    uid_list = list(people)
    events = defaultdict(list)
    event_sql = ",".join(f"'{e}'" for e in EVENTS)
    for t, t2 in m.chunks(m.BEH_START, m.OBS_END):
        rows = q(f"""
        SELECT device_id, TRY_CAST(ex_int_arg1 AS BIGINT), event_name, trigger, created_ts
        FROM immdata.user_behaviour
        WHERE created_ts >= '{t:%Y-%m-%d %H:%M:%S}'
          AND created_ts < '{t2:%Y-%m-%d %H:%M:%S}'
          AND event_name IN ({event_sql})
          AND (
            device_id IN ({m.sql_in(dev_ids)})
            OR TRY_CAST(ex_int_arg1 AS BIGINT) IN ({",".join(str(u) for u in uid_list)})
          )
        """)
        print(f"beh {t} {len(rows)}", flush=True)
        for device_id, uid, name, trigger, ts in rows:
            if uid and int(uid) in people:
                events[int(uid)].append((ts, name, trigger or ""))
            elif device_id in dev_ids:
                for u, p in people.items():
                    if device_id in p["devs"]:
                        events[u].append((ts, name, trigger or ""))

    if uid_list:
        pay_sql = " UNION ALL ".join(
            f"SELECT {u} AS userid, TIMESTAMP '{people[u]['pay_ts']:%Y-%m-%d %H:%M:%S}' AS pay_ts"
            for u in uid_list
        )
        tok_rows = q(f"""
        SELECT t.userid,
          SUM(COALESCE(t.inputusage,0)+COALESCE(t.outputusage,0)+COALESCE(t.cachedinputusage,0))
        FROM immdata.token_used_aw t
        INNER JOIN ({pay_sql}) p
          ON t.userid = p.userid AND t.timestamp < p.pay_ts
        WHERE t.timestamp >= '{m.ORDER_START:%Y-%m-%d %H:%M:%S}'
          AND t.timestamp < '{m.OBS_END:%Y-%m-%d %H:%M:%S}'
        GROUP BY t.userid
        """)
        tokens = {int(uid): int(n or 0) for uid, n in tok_rows}
    else:
        tokens = {}
    print("token users", len(tokens), flush=True)

    rows_out = []
    for uid, p in sorted(people.items(), key=lambda kv: kv[1]["pay_ts"]):
        ev = sorted(set(events.get(uid, [])))
        window = [
            e for e in ev
            if p["first_ts"] - timedelta(minutes=30) <= e[0] <= p["pay_ts"] + timedelta(minutes=1)
        ]
        cat, why = classify(window, p["pay_ts"])
        seq = collapse([f"{beijing(ts)} {label(name, trigger)}" for ts, name, trigger in window])
        flags = []
        blob = " ".join(label(n, t) for _ts, n, t in window)
        for key in ("领了赠送", "登录墙", "耗尽墙", "看视频字幕", "用AI字幕", "打开PDF", "切换翻译服务", "会员墙", "去定价", "点购买", "购买成功"):
            if key in blob:
                flags.append(key)
        hours = (p["pay_ts"] - p["first_ts"]).total_seconds() / 3600
        rows_out.append({
            "uid": uid,
            "ab": p["ab"],
            "pay": beijing(p["pay_ts"]),
            "hours": round(hours, 1),
            "tok": tokens.get(uid, 0),
            "kind": p["kind"],
            "otype": p["otype"],
            "amount": p["amount"] if p["amount"] is not None else "",
            "wall": "是" if p["wall"] else "否",
            "claim": "是" if p["claim"] else "否",
            "login": "是" if p["login"] else "否",
            "cat": cat,
            "why": why,
            "flags": "、".join(flags),
            "seq": seq[:3000],
        })

    wb = Workbook()
    ws = wb.active
    ws.title = "付款前行为"
    headers = [
        "账号", "字母", "付款时间(北京)", "激活后小时", "付款前token",
        "付费类型", "订单类型", "金额", "付款前耗尽墙", "领过赠送", "出现过登录墙",
        "归类", "归类依据", "行为", "付款前行为顺序",
    ]
    bold = Font(bold=True)
    for i, h in enumerate(headers, 1):
        ws.cell(1, i, h).font = bold
    for r, row in enumerate(rows_out, 2):
        vals = [
            row["uid"], row["ab"], row["pay"], row["hours"], row["tok"],
            row["kind"], row["otype"], row["amount"], row["wall"], row["claim"], row["login"],
            row["cat"], row["why"], row["flags"], row["seq"],
        ]
        for c, v in enumerate(vals, 1):
            ws.cell(r, c, v)
    sm = wb.create_sheet("汇总")
    sm["A1"] = "付费含兑换码，不含试用。截止北京 2026-10-16 01:05（订单日更水位）。新墙 f–j，剔老号。"
    sm["A1"].font = bold
    sm["A3"] = "范围"
    sm["B3"] = "人数"
    sm["A4"] = "付费账号"
    sm["B4"] = len(rows_out)
    sm["A5"] = "其中应用内付款"
    sm["B5"] = sum(1 for r in rows_out if r["kind"] == "应用内付款")
    sm["A6"] = "其中只有兑换码"
    sm["B6"] = sum(1 for r in rows_out if r["kind"] == "兑换码")
    sm["A7"] = "付款前已弹出耗尽墙"
    sm["B7"] = sum(1 for r in rows_out if r["wall"] == "是")
    sm["A8"] = "付款前没有耗尽墙"
    sm["B8"] = sum(1 for r in rows_out if r["wall"] == "否")
    sm["A10"] = "没有耗尽墙的人，按归类"
    sm["A10"].font = bold
    counts = defaultdict(int)
    for r in rows_out:
        if r["wall"] == "否":
            counts[r["cat"]] += 1
    rr = 11
    for cat, n in sorted(counts.items(), key=lambda kv: -kv[1]):
        sm.cell(rr, 1, cat)
        sm.cell(rr, 2, n)
        rr += 1
    sm.cell(rr + 1, 1, "付款前token是 token_used_aw 从北京 10-09 00:00 到这笔付费之前的 input+output+cached，不是赠送余额。金额为空的是兑换码。")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print("wrote", OUT, flush=True)
    print("no-wall categories", dict(counts), "sum", sum(counts.values()), flush=True)
    for row in rows_out:
        if row["wall"] == "否":
            print(row["cat"], row["uid"], row["kind"], row["why"], row["flags"], flush=True)


if __name__ == "__main__":
    main()
