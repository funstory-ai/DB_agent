"""Re-export task1 Excel: onetime excluded; all devices summed; flag shared devices."""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from db import query
from openpyxl import Workbook
from openpyxl.styles import Font

OUT = Path(__file__).resolve().parents[1] / "output"
OUT.mkdir(exist_ok=True)

DETAIL_SQL = """
WITH trial_30d AS (
  SELECT userid,
         MIN(createtime) AS first_trial_ts_30d,
         MAX(createtime) AS last_trial_ts_30d,
         COUNT(*) AS trial_orders_30d
  FROM immdata.imt_subscriptiongoodsorder
  WHERE COALESCE(isdeleted, 0) = 0
    AND orderstatus = 'finished'
    AND COALESCE(amountpaid, 0) = 0
    AND createtime >= DATEADD(day, -30, NOW())
    AND userid IS NOT NULL
    AND (
      subscriptiongoodssnap IS NULL
      OR CAST(subscriptiongoodssnap AS STRING) NOT LIKE '%\"SubscriptionType\":\"onetime%'
    )
  GROUP BY userid
),
paid_ever AS (
  -- 历史上现金付费 / onetime* / ordertype=code → 不算试用未付费（P024）
  SELECT DISTINCT userid
  FROM immdata.imt_subscriptiongoodsorder
  WHERE COALESCE(isdeleted, 0) = 0
    AND orderstatus = 'finished'
    AND userid IS NOT NULL
    AND (
      COALESCE(amountpaid, 0) > 0
      OR (
        COALESCE(subscriptionday, 0) > 0
        AND (
          CAST(subscriptiongoodssnap AS STRING) LIKE '%"SubscriptionType":"onetime%'
          OR LOWER(COALESCE(ordertype, '')) = 'code'
        )
      )
    )
),
cohort AS (
  SELECT t.*
  FROM trial_30d t
  LEFT JOIN paid_ever p ON t.userid = p.userid
  WHERE p.userid IS NULL
),
device_user_cnt AS (
  SELECT deviceid, COUNT(DISTINCT userid) AS users_on_device
  FROM immdata.imt_deviceinfo
  WHERE COALESCE(isdeleted, 0) = 0
    AND userid IS NOT NULL
    AND deviceid IS NOT NULL
    AND deviceid <> ''
  GROUP BY deviceid
),
user_devices AS (
  SELECT DISTINCT c.userid, d.deviceid
  FROM cohort c
  INNER JOIN immdata.imt_deviceinfo d
    ON c.userid = d.userid
   AND COALESCE(d.isdeleted, 0) = 0
   AND d.deviceid IS NOT NULL
   AND d.deviceid <> ''
),
user_device_stats AS (
  SELECT
    ud.userid,
    COUNT(DISTINCT ud.deviceid) AS device_cnt,
    COUNT(DISTINCT CASE WHEN duc.users_on_device > 1 THEN ud.deviceid END) AS shared_device_cnt,
    MAX(duc.users_on_device) AS max_users_on_any_device,
    MAX(CASE WHEN ud.deviceid = 'LK8j8qPLDa3gUgujz9Yt7RtmltkXP0bD' THEN 1 ELSE 0 END) AS has_device_LK8
  FROM user_devices ud
  INNER JOIN device_user_cnt duc ON ud.deviceid = duc.deviceid
  GROUP BY ud.userid
),
devices AS (
  SELECT DISTINCT deviceid FROM user_devices
),
beh AS (
  SELECT
    b.device_id,
    SUM(CASE WHEN b.event_name = 'translate_page' AND b.page_type = 'html' THEN 1 ELSE 0 END) AS web_cnt,
    SUM(CASE WHEN b.event_name = 'translate_video_subtitle' THEN 1 ELSE 0 END) AS video_cnt,
    SUM(CASE WHEN b.event_name IN ('translate_image', 'translate_image_client') THEN 1 ELSE 0 END) AS image_cnt,
    SUM(CASE WHEN b.event_name = 'translate_manga' THEN 1 ELSE 0 END) AS manga_cnt,
    SUM(CASE WHEN b.event_name = 'translate_page' AND b.page_type IN ('pdfReader', 'pdfProReader') THEN 1 ELSE 0 END) AS pdf_cnt
  FROM immdata.user_behaviour b
  INNER JOIN devices d ON b.device_id = d.deviceid
  WHERE b.created_ts >= DATEADD(day, -30, NOW())
    AND b.created_ts < NOW()
    AND (
      (b.event_name = 'translate_page' AND b.page_type IN ('html', 'pdfReader', 'pdfProReader'))
      OR b.event_name IN (
        'translate_video_subtitle',
        'translate_image',
        'translate_image_client',
        'translate_manga'
      )
    )
  GROUP BY b.device_id
),
user_agg AS (
  SELECT
    ud.userid,
    SUM(COALESCE(beh.web_cnt, 0)) AS web_translate_cnt,
    SUM(COALESCE(beh.video_cnt, 0)) AS video_translate_cnt,
    SUM(COALESCE(beh.image_cnt, 0)) AS image_translate_cnt,
    SUM(COALESCE(beh.manga_cnt, 0)) AS manga_translate_cnt,
    SUM(COALESCE(beh.pdf_cnt, 0)) AS pdf_translate_cnt
  FROM user_devices ud
  LEFT JOIN beh ON ud.deviceid = beh.device_id
  GROUP BY ud.userid
),
-- 账号级母语 P020：userid 下各设备按语言 SUM(cnt)，取最大
user_native_lang AS (
  SELECT
    ud.userid,
    nl.native_language,
    SUM(COALESCE(nl.cnt, 0)) AS cnt_sum
  FROM user_devices ud
  INNER JOIN immdata.imm_device_native_language nl
    ON ud.deviceid = nl.device_id
  WHERE nl.native_language IS NOT NULL
    AND nl.native_language <> ''
  GROUP BY ud.userid, nl.native_language
),
user_native_top AS (
  SELECT userid, native_language, cnt_sum
  FROM (
    SELECT
      userid,
      native_language,
      cnt_sum,
      ROW_NUMBER() OVER (
        PARTITION BY userid
        ORDER BY cnt_sum DESC, native_language
      ) AS rn
    FROM user_native_lang
  ) t
  WHERE rn = 1
)
SELECT
  c.userid,
  u.email,
  nt.native_language AS account_native_language,
  COALESCE(nt.cnt_sum, 0) AS native_language_cnt_sum,
  c.first_trial_ts_30d,
  c.last_trial_ts_30d,
  c.trial_orders_30d,
  COALESCE(ds.device_cnt, 0) AS device_cnt,
  COALESCE(ds.shared_device_cnt, 0) AS shared_device_cnt,
  COALESCE(ds.max_users_on_any_device, 0) AS max_users_on_any_device,
  COALESCE(ds.has_device_LK8, 0) AS has_device_LK8,
  CASE WHEN ds.userid IS NULL THEN 0 ELSE 1 END AS has_deviceinfo,
  COALESCE(ua.web_translate_cnt, 0) AS web_translate_cnt,
  COALESCE(ua.video_translate_cnt, 0) AS video_translate_cnt,
  COALESCE(ua.image_translate_cnt, 0) AS image_translate_cnt,
  COALESCE(ua.manga_translate_cnt, 0) AS manga_translate_cnt,
  COALESCE(ua.pdf_translate_cnt, 0) AS pdf_translate_cnt,
  COALESCE(ua.web_translate_cnt, 0)
    + COALESCE(ua.video_translate_cnt, 0)
    + COALESCE(ua.image_translate_cnt, 0)
    + COALESCE(ua.manga_translate_cnt, 0)
    + COALESCE(ua.pdf_translate_cnt, 0) AS translate_total_cnt
FROM cohort c
LEFT JOIN immdata.imt_userinfo u
  ON c.userid = u.id AND COALESCE(u.isdeleted, 0) = 0
LEFT JOIN user_native_top nt ON c.userid = nt.userid
LEFT JOIN user_device_stats ds ON c.userid = ds.userid
LEFT JOIN user_agg ua ON c.userid = ua.userid
ORDER BY translate_total_cnt DESC, c.userid
"""

HEADER_CN = [
    "用户ID",
    "邮箱",
    "账号母语",
    "母语合计cnt",
    "近30天首次试用时间",
    "近30天末次试用时间",
    "近30天试用订单数",
    "关联设备数",
    "共享设备数",
    "单设备最大绑定账号数",
    "是否绑定异常设备LK8",
    "是否有deviceinfo",
    "网页翻译次数",
    "视频翻译次数",
    "图片翻译次数",
    "漫画翻译次数",
    "PDF翻译次数",
    "五类合计",
]


def main():
    print("Querying...", flush=True)
    cols, rows = query(DETAIL_SQL)
    print(f"rows={len(rows)}", flush=True)

    wb = Workbook()
    note = wb.active
    note.title = "口径说明"
    for row in [
        ["项目", "说明"],
        ["邮箱", "imt_userinfo.email（id=userid，未软删）"],
        ["账号母语", "P020：该 userid 下全部设备在 imm_device_native_language 按语言 SUM(cnt)，取最大；母语合计cnt 为该合计数"],
        ["人群", "近30天 finished + amountpaid=0，排除 subscriptiongoodssnap.SubscriptionType 以 onetime 开头（淘宝兑换码一次性会员）"],
        ["未付费", "历史上无现金付费、无 onetime* 会员单（P024）"],
        ["翻译次数", "用户下全部设备累加（1–3 台正常）；不改为独占设备"],
        ["筛查字段", "共享设备数 / 单设备最大绑定账号数 / 是否绑定异常设备LK8，用于排查羊毛党，不改变次数口径"],
        ["异常设备LK8", "LK8j8qPLDa3gUgujz9Yt7RtmltkXP0bD：翻译极多且全库绑定约1690账号，导致多人次数同值"],
        ["网页", "translate_page + page_type=html"],
        ["视频", "translate_video_subtitle"],
        ["图片", "translate_image + translate_image_client"],
        ["漫画", "translate_manga"],
        ["PDF", "translate_page + page_type in (pdfReader, pdfProReader)"],
    ]:
        note.append(row)
    for c in note[1]:
        c.font = Font(bold=True)

    ws = wb.create_sheet("试用未付费用户翻译", 1)
    ws.append(HEADER_CN)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(list(r))

    out = OUT / "task1_trial_unpaid_30d_translate_by_type.xlsx"
    wb.save(out)
    print(f"Wrote {out}", flush=True)
    print("TOP8:", flush=True)
    for r in rows[:8]:
        print(r, flush=True)

    from patch_add_primary_platform import patch_workbook, fetch_platforms

    print("patch primary platform...", flush=True)
    userids = sorted({int(r[0]) for r in rows if r[0] is not None})
    plat = fetch_platforms(userids, days=30)
    print(f"with_platform={len(plat)}", flush=True)
    patch_workbook(
        out,
        out,
        [("试用未付费用户翻译", 1, "账号母语")],
        plat,
    )


if __name__ == "__main__":
    main()
