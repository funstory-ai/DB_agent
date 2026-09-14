"""为 task1/task2 追加「主要平台」：先收窄设备，再一次扫描 behaviour。

列名：主要平台（业务也称所属平台）
窗：与任务一致，近 30 天五类翻译事件
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from db import query
from openpyxl import load_workbook
from openpyxl.styles import Font

OUT = Path(__file__).resolve().parents[1] / "output"
DEFAULT_TASK2 = OUT / "task2_churn30d_vs_trial_unpaid_translate.xlsx"
DEFAULT_TASK1 = OUT / "task1_trial_unpaid_30d_translate_by_type.xlsx"

PLATFORM_EXPR = """
CASE
  WHEN lower(COALESCE(b.os_name, '')) IN ('ios', 'ipados', 'android') THEN '移动端'
  WHEN b.os_name = 'macOS' THEN 'Mac'
  WHEN b.os_name IN ('Windows', 'Linux', 'Chrome OS') THEN 'PC'
  WHEN b.platform_type IN ('mobile', 'tablet') THEN '移动端'
  WHEN b.platform_type LIKE 'iPhone%' OR b.platform_type LIKE 'iPad%' THEN '移动端'
  WHEN b.platform_type = 'desktop' THEN 'PC'
  ELSE '其他/未知'
END
"""

COL_NAME = "主要平台"
COL_NOTE = (
    "近30天五类翻译事件；os_name/platform_type 归并为 PC/Mac/移动端/其他未知；"
    "取事件数最多；无设备或无行为则空"
)


def collect_userids(task2: Path, task1: Path | None) -> list[int]:
    ids: set[int] = set()
    if task2.exists():
        wb = load_workbook(task2, read_only=True, data_only=True)
        for sheet in ("流失会员", "试用未付费"):
            if sheet not in wb.sheetnames:
                continue
            for i, r in enumerate(wb[sheet].iter_rows(values_only=True)):
                if i == 0 or r[1] is None:
                    continue
                ids.add(int(r[1]))
        wb.close()
    if task1 is not None and task1.exists():
        wb = load_workbook(task1, read_only=True, data_only=True)
        name = (
            "试用未付费用户翻译"
            if "试用未付费用户翻译" in wb.sheetnames
            else wb.sheetnames[-1]
        )
        for i, r in enumerate(wb[name].iter_rows(values_only=True)):
            if i == 0 or r[0] is None:
                continue
            try:
                ids.add(int(r[0]))
            except Exception:
                pass
        wb.close()
    return sorted(ids)


def fetch_platforms(userids: list[int], days: int = 30) -> dict[int, str]:
    result: dict[int, dict[str, int]] = {u: {} for u in userids}

    def chunks(xs, n):
        for i in range(0, len(xs), n):
            yield xs[i : i + n]

    for part in chunks(userids, 8000):
        ids = ",".join(str(u) for u in part)
        print(f"chunk users={len(part)} ...", flush=True)
        sql = f"""
        WITH ud AS (
          SELECT DISTINCT userid, deviceid
          FROM immdata.imt_deviceinfo
          WHERE COALESCE(isdeleted,0)=0
            AND userid IN ({ids})
            AND deviceid IS NOT NULL AND deviceid <> ''
        ),
        devices AS (SELECT DISTINCT deviceid FROM ud),
        beh AS (
          SELECT
            b.device_id,
            {PLATFORM_EXPR} AS major_platform,
            COUNT(*) AS cnt
          FROM immdata.user_behaviour b
          INNER JOIN devices d ON b.device_id = d.deviceid
          WHERE b.created_ts >= DATEADD(day,-{days},NOW()) AND b.created_ts < NOW()
            AND (
              (b.event_name='translate_page' AND b.page_type IN ('html','pdfReader','pdfProReader'))
              OR b.event_name IN (
                'translate_video_subtitle','translate_image','translate_image_client','translate_manga'
              )
            )
          GROUP BY b.device_id, major_platform
        ),
        user_plat AS (
          SELECT
            ud.userid,
            beh.major_platform,
            SUM(beh.cnt) AS cnt
          FROM ud
          INNER JOIN beh ON ud.deviceid = beh.device_id
          GROUP BY ud.userid, beh.major_platform
        )
        SELECT userid, major_platform, cnt FROM user_plat
        """
        _, rows = query(sql)
        print(f"  rows={len(rows)}", flush=True)
        for uid, plat, cnt in rows:
            uid = int(uid)
            plat = str(plat)
            cnt = int(cnt)
            result.setdefault(uid, {})
            result[uid][plat] = result[uid].get(plat, 0) + cnt

    out: dict[int, str] = {}
    for uid, dist in result.items():
        if not dist:
            continue
        preferred = {k: v for k, v in dist.items() if k != "其他/未知"}
        use = preferred if preferred else dist
        plat = max(use.items(), key=lambda kv: (kv[1], kv[0]))[0]
        out[uid] = plat
    return out


def _insert_or_fill_platform(ws, uid_col: int, plat: dict[int, str], after_header: str | None):
    headers = [c.value for c in ws[1]]
    if COL_NAME in headers:
        col = headers.index(COL_NAME) + 1
    else:
        if after_header and after_header in headers:
            col = headers.index(after_header) + 2
            ws.insert_cols(col)
        else:
            col = len(headers) + 1
        cell = ws.cell(1, col, COL_NAME)
        cell.font = Font(bold=True)
    for r in range(2, ws.max_row + 1):
        uid = ws.cell(r, uid_col).value
        if uid is None:
            continue
        ws.cell(r, col, plat.get(int(uid)))


def patch_workbook(
    path: Path,
    out_path: Path,
    sheets: list[tuple[str, int, str | None]],
    plat: dict[int, str],
):
    wb = load_workbook(path)
    if "口径说明" in wb.sheetnames:
        notes = [c.value for c in wb["口径说明"]["A"] if c.value]
        if COL_NAME not in notes:
            wb["口径说明"].append([COL_NAME, COL_NOTE])
    for sheet_name, uid_col, after_header in sheets:
        _insert_or_fill_platform(wb[sheet_name], uid_col, plat, after_header)
    wb.save(out_path)
    print(f"Wrote {out_path}", flush=True)


def patch_task_files(
    task2: Path = DEFAULT_TASK2,
    task1: Path | None = DEFAULT_TASK1,
    days: int = 30,
    task2_out: Path | None = None,
    task1_out: Path | None = None,
) -> dict[int, str]:
    print("collect userids...", flush=True)
    userids = collect_userids(task2, task1)
    print(f"n={len(userids)}", flush=True)
    print(f"fetch platforms ({days}d)...", flush=True)
    plat = fetch_platforms(userids, days=days)
    print(f"with_platform={len(plat)}", flush=True)

    if task2.exists():
        dest2 = task2_out or task2
        try:
            patch_workbook(
                task2,
                dest2,
                [
                    ("流失会员", 2, "账号母语"),
                    ("试用未付费", 2, "账号母语"),
                ],
                plat,
            )
        except PermissionError:
            dest2 = task2.with_name(task2.stem + "_with_platform.xlsx")
            print(f"locked, write {dest2}", flush=True)
            patch_workbook(
                task2,
                dest2,
                [
                    ("流失会员", 2, "账号母语"),
                    ("试用未付费", 2, "账号母语"),
                ],
                plat,
            )
    if task1 is not None and task1.exists():
        wb = load_workbook(task1, read_only=True)
        sheet = (
            "试用未付费用户翻译"
            if "试用未付费用户翻译" in wb.sheetnames
            else wb.sheetnames[-1]
        )
        wb.close()
        dest1 = task1_out or task1
        try:
            patch_workbook(
                task1,
                dest1,
                [(sheet, 1, "账号母语")],
                plat,
            )
        except PermissionError:
            dest1 = task1.with_name(task1.stem + "_with_platform.xlsx")
            print(f"locked, write {dest1}", flush=True)
            patch_workbook(
                task1,
                dest1,
                [(sheet, 1, "账号母语")],
                plat,
            )
    return plat


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--task2", type=Path, default=DEFAULT_TASK2)
    p.add_argument("--task1", type=Path, default=DEFAULT_TASK1)
    p.add_argument("--task2-out", type=Path, default=None)
    p.add_argument("--task1-out", type=Path, default=None)
    p.add_argument("--days", type=int, default=30)
    p.add_argument("--skip-task1", action="store_true")
    args = p.parse_args()
    patch_task_files(
        task2=args.task2,
        task1=None if args.skip_task1 else args.task1,
        days=args.days,
        task2_out=args.task2_out,
        task1_out=args.task1_out,
    )


if __name__ == "__main__":
    main()
