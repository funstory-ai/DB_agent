"""Task1 近30天抽样：邮箱+试用未付费归属；只打印。"""
from __future__ import annotations

import random
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from db import query
from openpyxl import load_workbook

SRC = Path(__file__).resolve().parents[1] / "output" / "task1_trial_unpaid_30d_translate_by_type.xlsx"


def main() -> None:
    wb = load_workbook(SRC, read_only=True, data_only=True)
    ws = wb["试用未付费用户翻译"]
    rows = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 or r[0] is None:
            continue
        rows.append({"userid": int(r[0]), "email": r[1]})
    wb.close()

    sample = random.Random(20260814).sample(rows, min(20, len(rows)))
    ids = ",".join(str(x["userid"]) for x in sample)

    _, erows = query(
        f"SELECT id, email FROM immdata.imt_userinfo WHERE COALESCE(isdeleted,0)=0 AND id IN ({ids})"
    )
    em = {int(r[0]): r[1] for r in erows}

    _, prows = query(
        f"""
        SELECT DISTINCT userid FROM immdata.imt_subscriptiongoodsorder
        WHERE COALESCE(isdeleted,0)=0 AND orderstatus='finished' AND userid IN ({ids})
          AND (
            COALESCE(amountpaid,0)>0
            OR (
              COALESCE(subscriptionday,0)>0
              AND (
                CAST(subscriptiongoodssnap AS STRING) LIKE '%"SubscriptionType":"onetime%'
                OR LOWER(COALESCE(ordertype,''))='code'
              )
            )
          )
        """
    )
    paid = {int(r[0]) for r in prows}

    _, trows = query(
        f"""
        SELECT DISTINCT userid FROM immdata.imt_subscriptiongoodsorder
        WHERE COALESCE(isdeleted,0)=0 AND orderstatus='finished'
          AND COALESCE(amountpaid,0)=0
          AND createtime >= DATEADD(day,-30,NOW())
          AND userid IN ({ids})
          AND (
            subscriptiongoodssnap IS NULL
            OR CAST(subscriptiongoodssnap AS STRING) NOT LIKE '%"SubscriptionType":"onetime%'
          )
        """
    )
    trial = {int(r[0]) for r in trows}

    fail = 0
    for r in sample:
        u = r["userid"]
        e_ok = (em.get(u) or None) == (r["email"] or None) or (u not in em and not r["email"])
        cohort_ok = u in trial and u not in paid
        if not (e_ok and cohort_ok):
            fail += 1
            print(f"FAIL userid={u} email_ok={e_ok} cohort_ok={cohort_ok}", flush=True)
        else:
            print(f"OK userid={u}", flush=True)

    n = len(sample)
    print(
        f"【抽样复核】task1近30天试用未付费；核对=email+人群归属；"
        f"N={n} 一致={n-fail}/{n} 结论={'通过' if fail==0 else '不通过'}",
        flush=True,
    )


if __name__ == "__main__":
    main()
