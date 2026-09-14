"""抽样核邮箱+账号母语：从已交付 Excel 抽 N 人，窄查询对拍（P019）。

只打印结论供对话汇报，不写抽样文件。
"""
from __future__ import annotations

import random
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from db import query
from openpyxl import load_workbook

OUT = Path(__file__).resolve().parents[1] / "output"
SRC = OUT / "task1_trial_unpaid_90d_translate_by_type_v4.xlsx"
N = 20
SEED = 20260813


def load_rows(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["试用未付费用户翻译"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = rows[0]
    # 用户ID, 邮箱, 账号母语, 母语合计cnt, ...
    data = []
    for r in rows[1:]:
        if r[0] is None:
            continue
        data.append(
            {
                "userid": int(r[0]),
                "email": r[1],
                "native": r[2],
                "cnt_sum": int(r[3] or 0),
            }
        )
    return header, data


def audit_email(userids: list[int]) -> dict[int, str | None]:
    ids = ",".join(str(u) for u in userids)
    sql = f"""
    SELECT id, email
    FROM immdata.imt_userinfo
    WHERE COALESCE(isdeleted,0)=0 AND id IN ({ids})
    """
    _, rows = query(sql)
    return {int(r[0]): r[1] for r in rows}


def audit_native(userids: list[int]) -> dict[int, tuple[str | None, int]]:
    """独立路径：deviceinfo → native_language → SUM(cnt) by lang → top1."""
    ids = ",".join(str(u) for u in userids)
    sql = f"""
    WITH ud AS (
      SELECT DISTINCT userid, deviceid
      FROM immdata.imt_deviceinfo
      WHERE COALESCE(isdeleted,0)=0
        AND userid IN ({ids})
        AND deviceid IS NOT NULL AND deviceid <> ''
    ),
    per_lang AS (
      SELECT
        ud.userid,
        nl.native_language,
        SUM(COALESCE(nl.cnt,0)) AS cnt_sum
      FROM ud
      INNER JOIN immdata.imm_device_native_language nl
        ON ud.deviceid = nl.device_id
      WHERE nl.native_language IS NOT NULL AND nl.native_language <> ''
      GROUP BY ud.userid, nl.native_language
    ),
    ranked AS (
      SELECT
        userid,
        native_language,
        cnt_sum,
        ROW_NUMBER() OVER (
          PARTITION BY userid ORDER BY cnt_sum DESC, native_language
        ) AS rn
      FROM per_lang
    )
    SELECT userid, native_language, cnt_sum
    FROM ranked WHERE rn = 1
    """
    _, rows = query(sql)
    return {int(r[0]): (r[1], int(r[2])) for r in rows}


def main():
    import argparse

    p = argparse.ArgumentParser()
    p.add_argument("-n", type=int, default=N)
    p.add_argument("--seed", type=int, default=SEED)
    p.add_argument("--src", type=Path, default=SRC)
    args = p.parse_args()

    if not args.src.exists():
        raise SystemExit(f"missing {args.src}")

    _, data = load_rows(args.src)
    rng = random.Random(args.seed)
    sample = rng.sample(data, min(args.n, len(data)))
    userids = [r["userid"] for r in sample]
    print(f"sampled={len(sample)} from {args.src.name}", flush=True)

    print("audit email...", flush=True)
    email_map = audit_email(userids)
    print("audit native (P020)...", flush=True)
    native_map = audit_native(userids)

    mismatch = 0
    results = []
    for r in sample:
        uid = r["userid"]
        e_ok = (email_map.get(uid) or None) == (r["email"] or None)
        # both None email: treat as match if key missing and report None
        if uid not in email_map and r["email"] in (None, ""):
            e_ok = True
        aud_n, aud_c = native_map.get(uid, (None, 0))
        rep_n = r["native"] if r["native"] not in (None, "") else None
        n_ok = aud_n == rep_n and aud_c == r["cnt_sum"]
        # both missing native
        if uid not in native_map and rep_n is None and r["cnt_sum"] == 0:
            n_ok = True
        ok = e_ok and n_ok
        if not ok:
            mismatch += 1
        results.append((uid, r, email_map.get(uid), aud_n, aud_c, e_ok, n_ok, ok))
        print(
            f"  userid={uid} email_ok={e_ok} native_ok={n_ok} "
            f"rep={rep_n}/{r['cnt_sum']} aud={aud_n}/{aud_c}",
            flush=True,
        )

    print(
        f"【抽样复核】任务=task1邮箱/母语；核对字段=email+native；"
        f"N={args.n} 一致={args.n - mismatch}/{args.n} "
        f"结论={'通过' if mismatch == 0 else '不通过'}",
        flush=True,
    )
    if mismatch:
        print("不一致/异常：", flush=True)
        for uid, r, ae, an, ac, e_ok, n_ok, ok in results:
            if not ok:
                print(
                    f"  userid={uid} email_ok={e_ok} native_ok={n_ok} "
                    f"rep_email={r['email']!r} aud_email={ae!r} "
                    f"rep_native={r['native']}/{r['cnt_sum']} aud_native={an}/{ac}",
                    flush=True,
                )
    else:
        print("不一致/异常：无", flush=True)


if __name__ == "__main__":
    main()
