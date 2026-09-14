"""Task2 抽样：从交付 Excel 抽检邮箱/母语/人群归属；只打印不落文件。"""
from __future__ import annotations

import random
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from db import query
from openpyxl import load_workbook

OUT = Path(__file__).resolve().parents[1] / "output"
SRC = OUT / "task2_churn30d_vs_trial_unpaid_translate.xlsx"
N_PER = 10  # 每人群默认 10，合计 20
SEED = 20260814


def load_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    data = []
    for r in rows[1:]:
        if r[1] is None:
            continue
        data.append(
            {
                "cohort": r[0],
                "userid": int(r[1]),
                "email": r[2],
                "native": r[3],
                "cnt_sum": int(r[4] or 0),
                "last_cover_end": r[5],
                "days_unpaid": r[8],
            }
        )
    return data


def audit_email(userids: list[int]) -> dict:
    ids = ",".join(str(u) for u in userids)
    _, rows = query(
        f"""
        SELECT id, email FROM immdata.imt_userinfo
        WHERE COALESCE(isdeleted,0)=0 AND id IN ({ids})
        """
    )
    return {int(r[0]): r[1] for r in rows}


def audit_native(userids: list[int]) -> dict:
    ids = ",".join(str(u) for u in userids)
    _, rows = query(
        f"""
        WITH ud AS (
          SELECT DISTINCT userid, deviceid
          FROM immdata.imt_deviceinfo
          WHERE COALESCE(isdeleted,0)=0 AND userid IN ({ids})
            AND deviceid IS NOT NULL AND deviceid <> ''
        ),
        per_lang AS (
          SELECT ud.userid, nl.native_language, SUM(COALESCE(nl.cnt,0)) AS cnt_sum
          FROM ud
          INNER JOIN immdata.imm_device_native_language nl ON ud.deviceid = nl.device_id
          WHERE nl.native_language IS NOT NULL AND nl.native_language <> ''
          GROUP BY ud.userid, nl.native_language
        ),
        ranked AS (
          SELECT userid, native_language, cnt_sum,
                 ROW_NUMBER() OVER (PARTITION BY userid ORDER BY cnt_sum DESC, native_language) rn
          FROM per_lang
        )
        SELECT userid, native_language, cnt_sum FROM ranked WHERE rn = 1
        """
    )
    return {int(r[0]): (r[1], int(r[2])) for r in rows}


def audit_churn_ok(userids: list[int]) -> dict:
    """确认：会员覆盖(含 onetime / ordertype=code) last_cover_end in [now-30, now)。"""
    ids = ",".join(str(u) for u in userids)
    _, rows = query(
        f"""
        SELECT userid,
               MAX(DATEADD(day, COALESCE(subscriptionday,0), createtime)) AS last_cover_end
        FROM immdata.imt_subscriptiongoodsorder
        WHERE COALESCE(isdeleted,0)=0 AND orderstatus='finished'
          AND COALESCE(subscriptionday,0)>0
          AND (
            COALESCE(amountpaid,0)>0
            OR CAST(subscriptiongoodssnap AS STRING) LIKE '%"SubscriptionType":"onetime%'
            OR LOWER(COALESCE(ordertype,''))='code'
          )
          AND userid IN ({ids})
        GROUP BY userid
        """
    )
    out = {}
    for r in rows:
        end = r[1]
        out[int(r[0])] = end
    return out


def audit_never_paid(userids: list[int]) -> set:
    """返回历史上有现金 / onetime / code 会员的 userid（试用侧不应出现）。"""
    ids = ",".join(str(u) for u in userids)
    _, rows = query(
        f"""
        SELECT DISTINCT userid FROM immdata.imt_subscriptiongoodsorder
        WHERE COALESCE(isdeleted,0)=0 AND orderstatus='finished'
          AND userid IN ({ids})
          AND (
            COALESCE(amountpaid,0)>0
            OR (
              COALESCE(subscriptionday,0)>0
              AND (
                CAST(subscriptiongoodssnap AS STRING) LIKE '%"SubscriptionType":"onetime%'
                OR LOWER(COALESCE(ordertype,''))='code'
              )
            )
          )
        """
    )
    return {int(r[0]) for r in rows}


def main():
    import argparse
    from datetime import datetime, timezone

    p = argparse.ArgumentParser()
    p.add_argument("--n-per", type=int, default=N_PER)
    p.add_argument("--seed", type=int, default=SEED)
    p.add_argument("--src", type=Path, default=SRC)
    args = p.parse_args()

    wb = load_workbook(args.src, read_only=True, data_only=True)
    churn = load_sheet(wb["流失会员"])
    trial = load_sheet(wb["试用未付费"])
    wb.close()

    rng = random.Random(args.seed)
    sample = rng.sample(churn, min(args.n_per, len(churn))) + rng.sample(
        trial, min(args.n_per, len(trial))
    )
    userids = [r["userid"] for r in sample]
    print(f"sampled churn={min(args.n_per,len(churn))} trial={min(args.n_per,len(trial))}", flush=True)

    email_map = audit_email(userids)
    native_map = audit_native(userids)
    churn_ends = audit_churn_ok([r["userid"] for r in sample if r["cohort"] == "流失会员"])
    paid_set = audit_never_paid([r["userid"] for r in sample if r["cohort"] == "试用未付费"])

    mismatch = 0
    now = datetime.now()
    for r in sample:
        uid = r["userid"]
        e_ok = (email_map.get(uid) or None) == (r["email"] or None)
        if uid not in email_map and r["email"] in (None, ""):
            e_ok = True
        aud_n, aud_c = native_map.get(uid, (None, 0))
        rep_n = r["native"] if r["native"] not in (None, "") else None
        n_ok = aud_n == rep_n and aud_c == r["cnt_sum"]
        if uid not in native_map and rep_n is None and r["cnt_sum"] == 0:
            n_ok = True

        cohort_ok = True
        if r["cohort"] == "流失会员":
            end = churn_ends.get(uid)
            if end is None:
                cohort_ok = False
            else:
                # 当日到期 DATEDIFF 可为 0；覆盖须落在近 30 天且已过期
                from datetime import timedelta

                in_30d = end >= (now - timedelta(days=30))
                days_ok = r["days_unpaid"] is None or int(r["days_unpaid"]) >= 0
                cohort_ok = end < now and in_30d and days_ok
        else:
            cohort_ok = uid not in paid_set

        ok = e_ok and n_ok and cohort_ok
        if not ok:
            mismatch += 1
            print(
                f"FAIL userid={uid} cohort={r['cohort']} email_ok={e_ok} native_ok={n_ok} cohort_ok={cohort_ok}",
                flush=True,
            )
        else:
            print(f"OK userid={uid} cohort={r['cohort']}", flush=True)

    n = len(sample)
    print(
        f"【抽样复核】任务=task2流失vs试用；核对=email+native+人群归属；"
        f"N={n} 一致={n - mismatch}/{n} 结论={'通过' if mismatch == 0 else '不通过'}",
        flush=True,
    )
    print("不一致/异常：" + ("无" if mismatch == 0 else "见上 FAIL"), flush=True)


if __name__ == "__main__":
    main()
